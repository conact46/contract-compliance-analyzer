"""
Contract Compliance Analyzer (Phase 1)
---------------------------------------
Research prototype for a master's thesis in Information Systems
(King Abdulaziz University). Screens business contract clauses against
six general Saudi laws using a knowledge graph, sentence embeddings
(semantic similarity) and fixed rule-based reasoning. No generative
language model is used in this phase; LLM integration is future work.

This file is a direct, cleaned-up port of the research notebook
`contract_analyzer_phase4.ipynb` for deployment as a Hugging Face Space,
so that domain experts can upload real contracts and evaluate the
system's outputs (correctness, clarity, traceability, usability) for the
thesis's expert-based evaluation chapter.
"""

import os
import re
import uuid
import time
import inspect
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape as _xml_escape
from datetime import datetime
from typing import Dict, List, Optional, Tuple
from dataclasses import dataclass, field
from enum import Enum

import numpy as np
import pandas as pd
import networkx as nx

import PyPDF2
import pdfplumber
from docx import Document as DocxDocument

from sentence_transformers import SentenceTransformer
import faiss
from sklearn.model_selection import train_test_split
from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, confusion_matrix
from sklearn.linear_model import LogisticRegression
from sklearn.ensemble import RandomForestClassifier, GradientBoostingClassifier
from sklearn.svm import SVC
from sklearn.preprocessing import LabelEncoder, StandardScaler

import plotly.graph_objects as go

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_RIGHT, TA_LEFT, TA_CENTER
from reportlab.platypus import HRFlowable
from reportlab.graphics.shapes import Drawing, Rect, Line, String
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import arabic_reshaper
from bidi.algorithm import get_display

import gradio as gr

# Register a bundled Arabic-capable font (Amiri, SIL OFL licensed, see
# fonts/Amiri-OFL-LICENSE.txt) so the PDF report can render Arabic text
# correctly instead of as unmapped glyph boxes.
_FONT_DIR = Path(__file__).parent / "fonts"
_ARABIC_FONT = "Amiri"
_ARABIC_FONT_BOLD = "Amiri-Bold"
try:
    pdfmetrics.registerFont(TTFont(_ARABIC_FONT, str(_FONT_DIR / "Amiri-Regular.ttf")))
    pdfmetrics.registerFont(TTFont(_ARABIC_FONT_BOLD, str(_FONT_DIR / "Amiri-Bold.ttf")))
    _ARABIC_FONT_OK = True
except Exception as e:
    print(f"Warning: could not load Amiri font ({e}); Arabic text in PDF reports may not render.")
    _ARABIC_FONT_OK = False


def _ar(text: str) -> str:
    """Reshapes and reorders Arabic text for correct right-to-left rendering
    in reportlab, which does not do this automatically. Safe to call on
    mixed Arabic/Latin/number text."""
    if not text:
        return text
    reshaped = arabic_reshaper.reshape(text)
    return get_display(reshaped)


_AR_CHAR_RE = re.compile(r'[؀-ۿ]')


def _detect_lang(text: str) -> str:
    """Cheap Arabic-vs-English detector used to keep clause-to-article
    matching within the same language. Each law now loads both Arabic and
    English illustrative articles side by side (see SAMPLE_LAWS below);
    without this, a law's semantic search could return an English article
    for an Arabic clause (or vice versa) purely on embedding similarity,
    which silently zeroes out keyword matching since the clause's literal
    words can never appear in a different-language article's keyword list."""
    if not text:
        return 'en'
    # Ratio over letters only, so that lines made mostly of underscores,
    # digits or punctuation (form fields such as "Date: ______") are judged
    # by their words.
    ar_chars = len(_AR_CHAR_RE.findall(text))
    latin = len(re.findall(r'[A-Za-z]', text))
    letters = ar_chars + latin
    if letters == 0:
        return 'en'
    return 'ar' if ar_chars > letters * 0.3 else 'en'


_AR_DIACRITICS_RE = re.compile(r'[ؐ-ًؚ-ٰٟۖ-ۭ]')
_AR_LETTER_MAP = str.maketrans({
    'أ': 'ا', 'إ': 'ا', 'آ': 'ا', 'ٱ': 'ا',
    'ى': 'ي', 'ة': 'ه', 'ؤ': 'و', 'ئ': 'ي',
    'ـ': '',  # tatweel
})


def normalize_text(text: str) -> str:
    """Normalizes text before literal term matching: lower-cases Latin
    letters and, for Arabic, removes diacritics and tatweel and unifies
    alef/hamza, alef maqsura and ta marbuta variants. Without this, a
    clause that writes a required term with a different (equally correct)
    spelling -- e.g. "إلغاء" vs "الغاء", or with tashkeel -- was counted as
    missing the term. Applied identically to the clause and to the
    article's required terms."""
    if not text:
        return ''
    text = _AR_DIACRITICS_RE.sub('', text)
    text = text.translate(_AR_LETTER_MAP)
    text = re.sub(r'\s+', ' ', text)
    return text.lower().strip()


_AR_DEFINITE_RE = re.compile(r'(?<!\S)(?:[وفبك]?ال|لل)(?=\S{2,})')


def term_form(text: str) -> str:
    """Form used for required-term matching: normalize_text plus removal of
    the Arabic definite article and its common attached prefixes (و/ف/ب/ك +
    ال, لل) at the start of each word, applied identically to the clause and
    to the term, so "السجل التجاري" matches a clause that writes
    "سجل تجاري" and vice versa."""
    return _AR_DEFINITE_RE.sub('', normalize_text(text))


def _unique(items):
    """Order-preserving de-duplication. Replaces the earlier list(set(...))
    calls, whose order depends on per-process hash seeds, so the same
    contract could show a different subset of findings on different runs."""
    return list(dict.fromkeys(items))


# Relevance gate (heuristic, documented in the thesis): a clause is treated
# as NOT APPLICABLE to a law -- and excluded from that law's compliance and
# risk figures -- when it is semantically far from every article of that law
# AND shares none of their required terms. Without this, a clause about
# payment was scored "non-compliant" with the Personal Data Protection Law
# simply for not mentioning consent.
#
# Set from the deployed model's observed similarities on the Arabic example
# contract, with the verified 78-article knowledge base: clauses that relate
# to a law scored 52-77% against its closest article; unrelated clauses
# scored 21-42% (the highest being an employee-obligations clause matched to
# PDPL Art. 8, and a confidentiality clause matched to Labor Law Art. 51).
# 0.47 sits midway between the two groups. A single shared generic term must
# not by itself make an unrelated clause count, so the term override requires
# at least half of the matched articles' required terms.
RELEVANCE_MIN_SIM = 0.47
RELEVANCE_TERM_OVERRIDE = 0.5

# Subject gate: a clause is screened against a law only if it mentions the
# subject that law regulates -- personal data for the PDPL, workers and wages
# for the Labor Law, consumers and online trade for the E-Commerce Law, and
# so on. Semantic similarity alone relates general contract wording to
# articles of laws that do not govern the contract: in a government supply
# contract "Contract copies" was matched to Labor Law Art. 51 (written
# employment contract) and "Taxes and fees" to PDPL Art. 8 (selection of
# processors), and in a business-to-business services contract "Parties and
# scope" was matched to E-Commerce Law Art. 13 (consumer's right to
# rescind). Generic contract terms that some articles require ("obligation",
# "contracting") made these look partly supported. A clause that mentions
# none of a law's subject words is therefore "not applicable" to that law,
# unless it is very close in meaning to one of its articles
# (SUBJECT_SIM_OVERRIDE). Words are matched as whole words after
# normalization, allowing a one-letter Arabic prefix and short inflectional
# suffixes (so "موظف" matches "والموظفين", but "اجر" does not match
# "اجراءات").
SUBJECT_SIM_OVERRIDE = 0.80
LAW_SUBJECTS = {
    "PDPL": {
        "ar": ["بيانات", "معلومات", "خصوصية"],
        "en": ["data", "personal information", "privacy"],
    },
    "LABOR": {
        "ar": ["عامل", "عمال", "عمالة", "موظف", "صاحب عمل", "عقد عمل", "أجور",
               "راتب", "رواتب", "توظيف", "استقدام", "سعودة", "توطين"],
        "en": ["employee", "worker", "staff", "personnel", "employer", "employment",
               "wage", "salary", "salaries", "labor", "labour", "workforce", "recruit"],
    },
    "ECC": {
        "ar": ["مستهلك", "تجارة إلكترونية", "متجر", "موقع إلكتروني", "منصة إلكترونية",
               "عبر الإنترنت", "الإنترنت"],
        "en": ["consumer", "e-commerce", "ecommerce", "electronic commerce", "online",
               "website", "web store", "internet"],
    },
    "CYBER": {
        "ar": ["معلوماتي", "معلوماتية", "شبكة", "اختراق", "دخول غير مشروع", "سيبراني",
               "أمن المعلومات", "نظام معلومات", "أنظمة المعلومات", "حاسب", "حاسوب",
               "برمجيات", "تنصت", "فيروس"],
        "en": ["cyber", "network", "hack", "unauthorized access", "unauthorised access",
               "information system", "information security", "computer", "malware",
               "virus", "intercept", "software"],
    },
    "COMMERCIAL": {
        "ar": ["شركة", "شركات", "سجل تجاري", "شريك", "شركاء", "مساهم", "مساهمة",
               "رأس مال", "عقد تأسيس", "نظام أساس"],
        "en": ["company", "companies", "commercial registration", "commercial register",
               "partner", "partnership", "shareholder", "capital", "articles of association",
               "incorporat", "limited liability", "llc"],
    },
    "ANTI_FRAUD": {
        "ar": ["منتج", "منتجات", "سلعة", "سلع", "بضاعة", "بضائع", "أصناف", "صنف",
               "مواصفات", "غش", "مغشوش", "تقليد", "مقلد", "مطابقة"],
        "en": ["product", "goods", "merchandise", "specification", "fraud",
               "counterfeit", "conformity", "defective"],
    },
}

_WORD_RE = re.compile(r'[\w\-]+')
_AR_ONE_LETTER_PREFIX = 'وفبلك'


def _tokens(text: str) -> List[str]:
    return _WORD_RE.findall(term_form(text))


def _word_matches(token: str, word: str) -> bool:
    """True if a clause token is the subject word, optionally with a
    one-letter Arabic prefix and a short suffix (up to 3 letters)."""
    candidates = [token]
    if token[:1] in _AR_ONE_LETTER_PREFIX and len(token) > len(word):
        candidates.append(token[1:])
    return any(c.startswith(word) and len(c) - len(word) <= 3 for c in candidates)


def mentions_subject(clause_tokens: List[str], subject_words: List[str]) -> List[str]:
    """Subject words (single or multi-word) found in the clause tokens."""
    found = []
    for phrase in subject_words:
        parts = _tokens(phrase)
        if not parts:
            continue
        n = len(parts)
        for i in range(len(clause_tokens) - n + 1):
            if all(_word_matches(clause_tokens[i + k], parts[k]) for k in range(n)):
                found.append(phrase)
                break
    return found

# Verdict thresholds on the combined score (unchanged from the notebook).
COMPLIANT_MIN = 0.55
PARTIAL_MIN = 0.35


# ============================================================
# 1. Enums and data classes
# ============================================================

class RiskLevel(Enum):
    LOW = "low"
    MEDIUM = "medium"
    HIGH = "high"
    CRITICAL = "critical"

    @property
    def emoji(self):
        return {"low": "\U0001F7E2", "medium": "\U0001F7E1", "high": "\U0001F7E0", "critical": "\U0001F534"}[self.value]

    @property
    def arabic(self):
        return {"low": "منخفض", "medium": "متوسط", "high": "عالي", "critical": "حرج"}[self.value]


class ComplianceStatus(Enum):
    COMPLIANT = "compliant"
    PARTIAL = "partial"
    NON_COMPLIANT = "non_compliant"
    NOT_APPLICABLE = "not_applicable"

    @property
    def icon(self):
        return {"compliant": "✅", "partial": "⚠️", "non_compliant": "❌", "not_applicable": "➖"}[self.value]

    def label(self, lang: str = 'en') -> str:
        if lang == 'ar':
            return {"compliant": "متوافق", "partial": "متوافق جزئيًا",
                    "non_compliant": "غير متوافق", "not_applicable": "غير منطبق"}[self.value]
        return {"compliant": "Compliant", "partial": "Partial",
                "non_compliant": "Non-compliant", "not_applicable": "Not applicable"}[self.value]


LAWS = {
    "PDPL": {"ar": "نظام حماية البيانات الشخصية", "en": "Personal Data Protection Law", "weight": 1.2},
    "ECC": {"ar": "نظام التجارة الإلكترونية", "en": "E-Commerce Law", "weight": 1.0},
    "LABOR": {"ar": "نظام العمل", "en": "Labor Law", "weight": 1.1},
    "CYBER": {"ar": "نظام مكافحة الجرائم المعلوماتية", "en": "Anti-Cyber Crime Law", "weight": 1.3},
    "COMMERCIAL": {"ar": "نظام الشركات", "en": "Companies Law", "weight": 1.0},
    "ANTI_FRAUD": {"ar": "نظام مكافحة الغش التجاري", "en": "Anti-Commercial Fraud Law", "weight": 1.1},
}


@dataclass
class LawArticle:
    id: str
    law_id: str
    number: str
    title: str
    content: str
    keywords: List[str] = field(default_factory=list)
    related_articles: List[str] = field(default_factory=list)
    risk_weight: float = 1.0
    embedding: Optional[np.ndarray] = None


@dataclass
class ContractClause:
    id: str
    contract_id: str
    number: int
    title: str
    content: str
    clause_type: str
    keywords: List[str] = field(default_factory=list)
    entities: Dict = field(default_factory=dict)
    embedding: Optional[np.ndarray] = None
    label: Optional[str] = None
    risk_label: Optional[str] = None


@dataclass
class AnalysisResult:
    status: ComplianceStatus
    risk_score: float
    reasoning: str
    issues: List[str] = field(default_factory=list)
    recommendations: List[str] = field(default_factory=list)
    confidence: float = 0.7
    # Traceability fields: which article the clause was matched to and the
    # exact numbers behind the verdict, so an expert can check the reasoning.
    best_article: Optional["LawArticle"] = None
    best_sim: float = 0.0
    avg_sim: float = 0.0
    matched_terms: int = 0
    total_terms: int = 0
    combined: float = 0.0
    missing_terms: List[str] = field(default_factory=list)


@dataclass
class TraceRow:
    """One clause x one law: the evidence behind a single verdict."""
    clause_number: int
    clause_title: str
    clause_excerpt: str
    law_id: str
    article_number: str
    article_title: str
    best_sim: float
    avg_sim: float
    matched_terms: int
    total_terms: int
    combined: float
    status: ComplianceStatus
    missing_terms: List[str] = field(default_factory=list)


@dataclass
class LawScore:
    law_id: str
    compliance_pct: Optional[float]  # None when no clause is relevant to this law
    risk_score: Optional[float]
    risk_level: Optional[RiskLevel]
    total_checked: int
    compliant: int
    partial: int
    non_compliant: int
    reasoning: str
    findings: List[str] = field(default_factory=list)
    not_applicable: int = 0

    @property
    def relevant(self) -> int:
        return self.compliant + self.partial + self.non_compliant


@dataclass
class FullReport:
    contract_id: str
    contract_name: str
    timestamp: str
    total_clauses: int
    laws_checked: List[str]
    law_scores: Dict[str, LawScore]
    overall_compliance: float
    overall_risk: float
    overall_level: RiskLevel
    summary: str
    risks: List[str] = field(default_factory=list)
    recommendations: List[str] = field(default_factory=list)
    learned_keywords: List[str] = field(default_factory=list)
    trace: List[TraceRow] = field(default_factory=list)
    contract_type: str = ''
    scope_notes: List[str] = field(default_factory=list)
    language: str = 'en'  # 'ar' or 'en' -- the contract's own language, used
    # to render save_pdf_report (and the on-screen summary) as a single
    # consistent language rather than the earlier English/Arabic pairing.


@dataclass
class TrainingData:
    X_train: np.ndarray
    X_val: np.ndarray
    X_test: np.ndarray
    y_train: np.ndarray
    y_val: np.ndarray
    y_test: np.ndarray
    label_encoder: LabelEncoder
    scaler: StandardScaler


# ============================================================
# 2. Knowledge graphs
# ============================================================

# NOTE: switched from the notebook's original "all-MiniLM-L6-v2" (English-only;
# trained exclusively on English sentence pairs) to a multilingual model.
# The loaded law content (SAMPLE_LAWS) is Arabic, so an English-only encoder
# produced near-meaningless, non-discriminative embeddings for it -- semantic
# similarity (60% of the ReasoningEngine's combined score) was effectively
# noise, which meant clauses almost never crossed the 0.35/0.55 thresholds
# and the tool could not realistically ever report "Compliant" or "Partial".
# "paraphrase-multilingual-MiniLM-L12-v2" is trained on 50+ languages
# including Arabic, is the same sentence-transformers family/size class, and
# is a drop-in replacement -- the embedding dimension is read dynamically
# from the model everywhere in this file, so no other code needs to change.
EMBED_MODEL_NAME = "paraphrase-multilingual-MiniLM-L12-v2"


class LawKnowledgeGraph:
    """Knowledge graph of the loaded laws:

        Law --contains--> Article --requires_term--> Term

    Both steps of the analysis read from this graph: retrieval only considers
    the articles the graph links to the selected law (``articles_of``), and
    the term check uses the terms the graph links to each matched article
    (``required_terms``). FAISS provides the semantic ranking."""

    def __init__(self, embedder: SentenceTransformer):
        self.embedder = embedder
        self.graph = nx.DiGraph()
        self.articles: Dict[str, LawArticle] = {}
        self.index = None
        self.id_map: List[str] = []

        for law_id, info in LAWS.items():
            self.graph.add_node(law_id, node_type="law", **info)
            # Law --has_subject--> Subject: the words that show a clause
            # deals with what this law regulates (see LAW_SUBJECTS).
            for lang_words in LAW_SUBJECTS.get(law_id, {}).values():
                for word in lang_words:
                    sn = f"SUBJECT::{normalize_text(word)}"
                    self.graph.add_node(sn, node_type="subject", label=word)
                    self.graph.add_edge(law_id, sn, relation="has_subject")

    @staticmethod
    def _term_node(term: str) -> str:
        return f"TERM::{normalize_text(term)}"

    def add_article(self, article: LawArticle):
        text = f"{article.title} {article.content}"
        article.embedding = self.embedder.encode(text)

        self.articles[article.id] = article
        self.graph.add_node(article.id, node_type="article", law_id=article.law_id,
                            number=article.number, title=article.title,
                            lang=_detect_lang(text))
        self.graph.add_edge(article.law_id, article.id, relation="contains")

        # A required term may list spelling/inflection alternatives separated
        # by "|" (e.g. "ثلاثين|ثلاثون", "destruction|destroy"); the first is
        # the label shown to the user, and the term counts as present if any
        # alternative appears in the clause.
        for term in article.keywords:
            alts = [a.strip() for a in term.split('|') if a.strip()]
            tn = self._term_node(alts[0])
            if tn not in self.graph:
                self.graph.add_node(tn, node_type="term", label=alts[0],
                                    norms=[term_form(a) for a in alts])
            else:
                known = self.graph.nodes[tn]["norms"]
                known.extend(n for n in (term_form(a) for a in alts) if n not in known)
            self.graph.add_edge(article.id, tn, relation="requires_term")

        self.index = None

    def articles_of(self, law_id: str) -> List[str]:
        """Article ids the graph links to a law (Law --contains--> Article)."""
        if law_id not in self.graph:
            return []
        return [n for n in self.graph.successors(law_id)
                if self.graph.edges[law_id, n].get("relation") == "contains"
                and self.graph.nodes[n].get("node_type") == "article"]

    def subject_words(self, law_id: str) -> List[str]:
        """Words the graph links to a law (Law --has_subject--> Subject)."""
        if law_id not in self.graph:
            return []
        return [self.graph.nodes[n]["label"] for n in self.graph.successors(law_id)
                if self.graph.edges[law_id, n].get("relation") == "has_subject"]

    def required_terms(self, article_id: str) -> List[Tuple[str, List[str]]]:
        """(display label, normalized alternatives) of each term an article
        requires (Article --requires_term--> Term)."""
        if article_id not in self.graph:
            return []
        return [(self.graph.nodes[n]["label"], self.graph.nodes[n]["norms"])
                for n in self.graph.successors(article_id)
                if self.graph.edges[article_id, n].get("relation") == "requires_term"]

    def build_index(self):
        if not self.articles:
            return
        embeddings = []
        self.id_map = []
        for aid, art in self.articles.items():
            if art.embedding is not None:
                embeddings.append(art.embedding)
                self.id_map.append(aid)
        if embeddings:
            arr = np.array(embeddings).astype('float32')
            faiss.normalize_L2(arr)
            self.index = faiss.IndexFlatIP(arr.shape[1])
            self.index.add(arr)

    def search(self, query: str, top_k: int = 5, law_filter: List[str] = None,
               query_emb: Optional[np.ndarray] = None) -> List[Tuple[LawArticle, float]]:
        """Returns up to top_k (article, cosine similarity) pairs, best first.

        FAISS ranks every article (an exhaustive search -- the corpus is tens
        of articles, so this is cheap); the result is then restricted to the
        articles the knowledge graph links to the requested law(s). Ranking
        the whole corpus before filtering avoids an earlier bug where a
        global top-k cut could drop the requested law's own articles."""
        if self.index is None:
            self.build_index()
        if self.index is None:
            return []

        if query_emb is None:
            query_emb = self.embedder.encode([query])[0]
        q = np.array([query_emb], dtype='float32')
        faiss.normalize_L2(q)
        sims, idxs = self.index.search(q, len(self.id_map))

        allowed = None
        if law_filter:
            allowed = {aid for law in law_filter for aid in self.articles_of(law)}

        scored: List[Tuple[LawArticle, float]] = []
        for sim, i in zip(sims[0], idxs[0]):
            if i < 0:
                continue
            aid = self.id_map[i]
            if allowed is not None and aid not in allowed:
                continue
            scored.append((self.articles[aid], float(sim)))

        # Prefer same-language articles: term matching is literal-text
        # based, so a different-language article can never score on that
        # half of the combined score. Fall back to the other language only
        # when the law has no article at all in the clause's language.
        query_lang = _detect_lang(query)
        same_lang = [p for p in scored if self.graph.nodes[p[0].id].get("lang") == query_lang]
        if same_lang:
            return same_lang[:top_k]
        return scored[:top_k]


class ContractKnowledgeGraph:
    """Knowledge graph of ONE contract while it is being analyzed:

        Contract --contains--> Clause --matched_to--> Article (law graph)

    A new instance is created for every analysis request and discarded when
    the request ends, so no contract text outlives the request that
    produced its report (see the privacy note in the interface). The
    clause-to-article edges carry the evidence behind each verdict and are
    what the traceability table is built from."""

    def __init__(self, embedder: SentenceTransformer):
        self.embedder = embedder
        self.graph = nx.DiGraph()
        self.contracts: Dict[str, Dict] = {}
        self.clauses: Dict[str, ContractClause] = {}
        self.learned_patterns: Dict[str, List[str]] = {}

    def add_contract(self, contract_id: str, name: str, metadata: Dict = None):
        self.contracts[contract_id] = {"name": name, **(metadata or {})}
        self.graph.add_node(contract_id, node_type="contract", name=name)

    def add_clause(self, clause: ContractClause):
        text = f"{clause.title} {clause.content}"
        clause.embedding = self.embedder.encode(text)

        self.clauses[clause.id] = clause
        self.graph.add_node(clause.id, node_type="clause", title=clause.title, number=clause.number)
        self.graph.add_edge(clause.contract_id, clause.id, relation="contains")

        ctype = clause.clause_type
        if ctype not in self.learned_patterns:
            self.learned_patterns[ctype] = []
        for kw in clause.keywords:
            if kw not in self.learned_patterns[ctype]:
                self.learned_patterns[ctype].append(kw)

    def add_match(self, clause_id: str, law_id: str, result: "AnalysisResult"):
        """Records the clause's verdict against one law as an edge to the
        best-matching article, keyed by law so one clause can link to
        several laws' articles."""
        art = result.best_article
        target = art.id if art is not None else f"NO_ARTICLE::{law_id}"
        if target not in self.graph:
            self.graph.add_node(target, node_type="article_ref")
        self.graph.add_edge(clause_id, target, relation="matched_to",
                            law_id=law_id, result=result)

    def get_clauses(self, contract_id: str) -> List[ContractClause]:
        if contract_id not in self.graph:
            return []
        ids = [n for n in self.graph.successors(contract_id)
               if self.graph.nodes[n].get("node_type") == "clause"]
        return sorted((self.clauses[i] for i in ids), key=lambda c: c.number)

    def trace(self, contract_id: str, law_order: List[str]) -> List[TraceRow]:
        """Builds the clause-by-clause evidence table from the graph."""
        rows: List[TraceRow] = []
        for clause in self.get_clauses(contract_id):
            edges = {d["law_id"]: d["result"]
                     for _, _, d in self.graph.out_edges(clause.id, data=True)
                     if d.get("relation") == "matched_to"}
            for law_id in law_order:
                r = edges.get(law_id)
                if r is None:
                    continue
                excerpt = clause.content.strip()
                if len(excerpt) > 90:
                    excerpt = excerpt[:87] + "…"
                rows.append(TraceRow(
                    clause_number=clause.number, clause_title=clause.title,
                    clause_excerpt=excerpt, law_id=law_id,
                    article_number=r.best_article.number if r.best_article else "—",
                    article_title=r.best_article.title if r.best_article else "—",
                    best_sim=r.best_sim, avg_sim=r.avg_sim,
                    matched_terms=r.matched_terms, total_terms=r.total_terms,
                    combined=r.combined, status=r.status,
                    missing_terms=list(r.missing_terms),
                ))
        return rows


# ============================================================
# 3. Contract parsing
# ============================================================

class ContractParser:
    """Extracts clauses from a contract (Arabic or English)."""

    CLAUSE_PATTERNS = [
        r'(?:المادة|البند|الفقرة)\s*(?:رقم)?\s*[:\-]?\s*(\d+)',
        r'(?:Article|Clause|Section)\s*[:\-]?\s*(\d+)',
        r'^(\d+)[.\-\)]\s+',
    ]

    CLAUSE_TYPES = {
        'privacy': ['خصوصية', 'بيانات', 'سرية', 'موافقة', 'شخصية'],
        'payment': ['دفع', 'سداد', 'مقابل', 'مالي', 'ريال'],
        'liability': ['مسؤولية', 'ضمان', 'تعويض', 'إخلاء'],
        'termination': ['إنهاء', 'فسخ', 'انتهاء', 'إلغاء'],
        'security': ['أمن', 'حماية', 'اختراق', 'سيبراني'],
        'dispute': ['نزاع', 'تحكيم', 'محكمة', 'خلاف'],
        'compliance': ['امتثال', 'التزام', 'قانون', 'لائحة'],
    }

    def parse_text(self, text: str) -> List[Dict]:
        clauses = []
        lines = text.split('\n')
        current = None
        num = 0
        seen_first_line = False

        for line in lines:
            line = line.strip()
            if not line:
                continue
            # Table-of-contents entries ("9 السرية ...........14") repeat
            # the headings and are not clause text.
            if re.search(r'\.{5,}|…{3,}', line):
                continue

            is_header = any(re.match(p, line) for p in self.CLAUSE_PATTERNS)
            # A numbered heading written without a delimiter, as in
            # official templates ("24 الالتزامات العامة").
            is_header = is_header or bool(
                len(line) <= 60 and re.match(r'^\d{1,3}\s+[^\d\s]', line))
            # A short "Title: text" line starts a clause; a line that ends
            # with a colon introduces a list ("... الآتي:", "as follows:")
            # and belongs to the clause it is in.
            is_header = is_header or (len(line) < 80 and ':' in line
                                      and not line.rstrip().endswith(':'))

            if is_header:
                seen_first_line = True
                if current and current['content'].strip():
                    clauses.append(current)
                num += 1
                # A numbered line that is long is a clause whose text sits on
                # the same line ("1. The Supplier shall ..."); keep that text
                # as the clause content instead of dropping the clause for
                # having an "empty" body.
                inline = line if len(line) > 80 else ''
                current = {
                    'number': num,
                    'title': line if len(line) <= 100 else line[:97] + '…',
                    'content': inline,
                    'type': self._detect_type(line)
                }
            elif current:
                current['content'] += ' ' + line
            else:
                # Text before the first clause header. The very first short
                # line is the document's title ("عقد تقديم خدمات تقنية"), not
                # a clause -- it was previously counted as an extra clause.
                if not seen_first_line and len(line) <= 100 and not line.endswith(('.', '،', '؛')):
                    seen_first_line = True
                    continue
                seen_first_line = True
                num += 1
                label = 'تمهيد' if _detect_lang(line) == 'ar' else 'Preamble'
                current = {'number': num, 'title': label, 'content': line, 'type': 'general'}

        if current and current['content'].strip():
            clauses.append(current)
        return clauses

    def _detect_type(self, text: str) -> str:
        text_lower = text.lower()
        for ctype, keywords in self.CLAUSE_TYPES.items():
            if any(kw in text_lower for kw in keywords):
                return ctype
        return 'general'

    def extract_keywords(self, text: str) -> List[str]:
        words = re.findall(r'[\w؀-ۿ]{3,}', text)
        stopwords = {'من', 'إلى', 'على', 'في', 'أن', 'هذا', 'التي', 'الذي'}
        return _unique(w for w in words if w.lower() not in stopwords)[:15]

    @staticmethod
    def _arabic_order_score(text: str) -> int:
        """How much a text looks like Arabic in logical (reading) order.

        Some PDFs store Arabic glyphs in visual order, and an extractor then
        returns every Arabic word spelled backwards ("ةطخلا" for "الخطة").
        Words that begin with the definite article are frequent in logical
        order, and in reversed text they end with it instead, so the
        difference between the two counts tells the orders apart. Returns 0
        for text with no Arabic."""
        words = re.findall(r'[\u0600-\u06FF]{3,}', text)
        starts = sum(1 for w in words if w.startswith('ال'))
        ends = sum(1 for w in words if w.endswith('لا'))
        return starts - ends

    @staticmethod
    def _reverse_visual_line(line: str) -> str:
        """Converts one visually ordered line to logical order: reverses the
        characters, then restores runs of Latin letters and digits (which
        were stored left to right) and mirrors brackets."""
        rev = line[::-1]
        rev = re.sub(r'[A-Za-z0-9][A-Za-z0-9.,:/_\-]*[A-Za-z0-9]|[A-Za-z0-9]',
                     lambda m: m.group(0)[::-1], rev)
        return rev.translate(str.maketrans('()[]{}<>', ')(][}{><'))

    @staticmethod
    def _repair_ligatures(text: str) -> str:
        """Some PDFs store "ال" followed by a hamza-alef as one ligature
        glyph, which extractors return in the wrong order ("اإلصدار" for
        "الإصدار", "األمثلة" for "الأمثلة"). The broken sequences never
        occur in correct Arabic, so they are swapped back."""
        return re.sub(r'ا([أإآ])ل', r'ال\1', text)

    @staticmethod
    def _drop_page_furniture(pages: List[str]) -> List[str]:
        """Removes running headers and footers: lines that recur on many
        pages (ignoring digits, underscores and spacing, so "page 3 of 46"
        and "page 4 of 46" count as the same line). In a 46-page contract
        template the footer otherwise became 44 of its 101 "clauses"."""
        if len(pages) < 3:
            return pages
        def key(line: str) -> str:
            return re.sub(r'[\d_\s\.\-/|:]+', '', line)
        counts: Dict[str, int] = {}
        for page in pages:
            for k in {key(l) for l in page.split("\n") if key(l)}:
                counts[k] = counts.get(k, 0) + 1
        limit = max(3, int(0.5 * len(pages)))
        repeated = {k for k, n in counts.items() if n >= limit}
        return ["\n".join(l for l in page.split("\n") if key(l) not in repeated)
                for page in pages]

    @staticmethod
    def _without_phantom_spaces(page):
        """Word-generated Arabic PDFs can draw a space character inside the
        wide final-form glyph of a word's last letter (ة, ت). The extractor
        treats it as a word break and returns "المعلوما ت" for "المعلومات".
        A space whose centre lies inside another glyph on the same line is
        not a real gap, so it is dropped before the text is extracted."""
        glyphs = [c for c in page.chars if c.get("text", "").strip()]
        def phantom(c) -> bool:
            mid = (c["x0"] + c["x1"]) / 2
            return any(g["x0"] + 0.5 < mid < g["x1"] - 0.5
                       and abs(g["top"] - c["top"]) < 0.5 * max(g.get("size", 1), 1)
                       for g in glyphs)
        def keep(obj) -> bool:
            if obj.get("object_type") != "char" or obj.get("text", "").strip():
                return True
            return not phantom(obj)
        try:
            return page.filter(keep)
        except Exception:
            return page

    def parse_pdf(self, path: str) -> Tuple[str, List[Dict]]:
        """Extracts the text of a PDF with pdfplumber and with PyPDF2 and
        keeps the version whose Arabic reads in logical order. pdfplumber
        keeps the layout better and is preferred for English; for Arabic
        PDFs stored in visual order it returns reversed words, which made
        every clause unmatchable, so PyPDF2's text, or pdfplumber's text
        with each line reversed, is used instead. Running headers and
        footers are then removed and broken lam-alef ligatures repaired."""
        candidates: List[List[str]] = []
        try:
            with pdfplumber.open(path) as pdf:
                plumber, flipped = [], []
                for page in pdf.pages:
                    page = self._without_phantom_spaces(page)
                    plumber.append(page.extract_text() or "")
                    # A lam-alef ligature is one glyph whose text is already
                    # in logical order ("لا"). Reversing a visual-order line
                    # character by character would turn it into "ال"
                    # ("السالمة" for "السلامة"), so the glyph's own text is
                    # flipped first and the line reversal restores it.
                    chars = [dict(c, text=c["text"][::-1])
                             if re.fullmatch(r'[؀-ۿ]{2,3}', c.get("text", ""))
                             else c
                             for c in page.chars]
                    flipped.append(pdfplumber.utils.extract_text(chars) or "")
            candidates.append(plumber)
            candidates.append(["\n".join(self._reverse_visual_line(l) for l in p.split("\n"))
                               for p in flipped])
        except Exception as e:
            print(f"pdfplumber failed ({e})")
        try:
            with open(path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                candidates.append([(page.extract_text() or "") for page in reader.pages])
        except Exception as e:
            print(f"PyPDF2 failed ({e})")
        candidates = [c for c in candidates if "".join(c).strip()]
        if not candidates:
            return "", []
        # max() keeps the first candidate on ties, so pdfplumber wins for
        # English text and for Arabic text that is already in logical order.
        pages = max(candidates, key=lambda c: self._arabic_order_score("\n".join(c)))
        pages = self._drop_page_furniture(pages)
        text = self._repair_ligatures("\n".join(pages))
        return text, self.parse_text(text)

    def parse_docx(self, path: str) -> Tuple[str, List[Dict]]:
        doc = DocxDocument(path)
        parts = [p.text for p in doc.paragraphs]
        # Contracts often put clauses inside tables; include table cell text.
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        text = "\n".join(parts)
        return text, self.parse_text(text)


# ============================================================
# 4. Reasoning engine (fixed rule-based formula)
# ============================================================

class ReasoningEngine:
    def __init__(self, embedder: SentenceTransformer):
        self.embedder = embedder

    def analyze(self, clause: ContractClause, matches: List[Tuple[LawArticle, float]],
                law_id: str, law_kg: LawKnowledgeGraph, lang: str = 'en') -> AnalysisResult:
        """Scores one clause against the articles retrieved for one law.

        combined = 0.4 x (share of the matched articles' required terms that
        appear in the clause) + 0.6 x (mean cosine similarity to those
        articles). Verdict: > 0.55 compliant, > 0.35 partial, otherwise
        non-compliant -- unless the relevance gate marks the clause as not
        applicable to this law."""
        law_ar = LAWS.get(law_id, {}).get('ar', law_id)
        law_en = LAWS.get(law_id, {}).get('en', law_id)
        is_ar = (lang == 'ar')

        if not matches:
            return AnalysisResult(
                status=ComplianceStatus.NOT_APPLICABLE,
                risk_score=0.0,
                reasoning=(f"لا توجد مواد ذات صلة من {law_ar}" if is_ar
                           else f"No relevant articles from {law_en}"),
                confidence=0.5,
            )

        clause_norm = term_form(f"{clause.title} {clause.content}")

        total_kw, matched_kw = 0, 0
        missing_by_article: List[Tuple[LawArticle, List[str]]] = []
        for art, _ in matches:
            terms = law_kg.required_terms(art.id)
            total_kw += len(terms)
            missing = [label for label, norms in terms
                       if not any(n and n in clause_norm for n in norms)]
            matched_kw += len(terms) - len(missing)
            missing_by_article.append((art, missing))

        sims = [s for _, s in matches]
        best_art, best_sim = matches[0]
        avg_sim = float(np.mean(sims))
        kw_ratio = matched_kw / max(total_kw, 1)
        combined = 0.4 * kw_ratio + 0.6 * avg_sim
        best_missing = missing_by_article[0][1]

        trace_kwargs = dict(best_article=best_art, best_sim=best_sim, avg_sim=avg_sim,
                            matched_terms=matched_kw, total_terms=total_kw,
                            combined=combined, missing_terms=best_missing)

        subject_found = mentions_subject(_tokens(f"{clause.title} {clause.content}"),
                                         law_kg.subject_words(law_id))
        if not subject_found and best_sim < SUBJECT_SIM_OVERRIDE:
            return AnalysisResult(
                status=ComplianceStatus.NOT_APPLICABLE,
                risk_score=0.0,
                reasoning=(f"البند لا يتناول موضوع {law_ar}" if is_ar else
                           f"Clause does not deal with the subject of {law_en}"),
                confidence=0.5,
                **trace_kwargs,
            )

        if best_sim < RELEVANCE_MIN_SIM and kw_ratio < RELEVANCE_TERM_OVERRIDE:
            return AnalysisResult(
                status=ComplianceStatus.NOT_APPLICABLE,
                risk_score=0.0,
                reasoning=(f"البند لا يتعلق بـ{law_ar} (أعلى تشابه {best_sim*100:.0f}%)"
                           if is_ar else
                           f"Clause not related to {law_en} (best similarity {best_sim*100:.0f}%)"),
                confidence=0.5,
                **trace_kwargs,
            )

        if combined > COMPLIANT_MIN:
            status = ComplianceStatus.COMPLIANT
            risk = 15 + (1 - combined) * 20
        elif combined > PARTIAL_MIN:
            status = ComplianceStatus.PARTIAL
            risk = 40 + (1 - combined) * 25
        else:
            status = ComplianceStatus.NON_COMPLIANT
            risk = 65 + (1 - combined) * 30

        issues, recommendations = [], []
        if status != ComplianceStatus.COMPLIANT:
            for art, missing in missing_by_article:
                if not missing:
                    continue
                missing_str = ', '.join(missing[:2])
                if is_ar:
                    issues.append(f"البند {clause.number} مقابل المادة {art.number} ({law_ar}): مصطلحات غير مذكورة ({missing_str})")
                    recommendations.append(f"مراجعة تغطية: {missing_str}")
                else:
                    issues.append(f"Clause {clause.number} vs Art. {art.number} ({law_en}): terms not mentioned ({missing_str})")
                    recommendations.append(f"Review coverage of: {missing_str}")

        reasoning = (
            f"تحليل ضد {law_ar}: تشابه {avg_sim*100:.0f}%، مصطلحات {kw_ratio*100:.0f}%، مجمع {combined*100:.0f}%"
            if is_ar else
            f"Analysis vs {law_en}: similarity {avg_sim*100:.0f}%, terms {kw_ratio*100:.0f}%, "
            f"combined {combined*100:.0f}%"
        )

        return AnalysisResult(
            status=status,
            risk_score=min(100, max(0, risk)),
            reasoning=reasoning,
            issues=issues[:5],
            recommendations=_unique(recommendations)[:5],
            confidence=0.6 + 0.3 * combined,
            **trace_kwargs,
        )

    def generate_summary(self, overall_risk: Optional[float], name: str, lang: str = 'en',
                         n_applicable: int = 0, n_checked: int = 0) -> str:
        if overall_risk is None:
            if lang == 'ar':
                return (f"العقد: {name}\nالحالة: لم يُعثر على بنود تتعلق بالأنظمة المختارة\n"
                        f"التوصية: جرّب/ي أنظمة أخرى أو راجع/ي جدول تتبع البنود")
            return (f"Contract: {name}\nStatus: no clause was found relevant to the selected laws\n"
                    f"Recommendation: try other laws or check the clause traceability table")
        if overall_risk < 30:
            status_en, status_ar = "Contract looks good", "العقد جيد"
            action_en, action_ar = "Can proceed", "يمكن المتابعة"
        elif overall_risk < 50:
            status_en, status_ar = "Needs improvements", "يحتاج تحسينات"
            action_en, action_ar = "Review flagged clauses", "مراجعة البنود المحددة"
        elif overall_risk < 75:
            status_en, status_ar = "Notable risk", "مخاطر ملحوظة"
            action_en, action_ar = "Legal consultation advised", "استشارة قانونية"
        else:
            status_en, status_ar = "High risk", "مخاطر عالية"
            action_en, action_ar = "Substantial revisions required", "تعديلات جوهرية مطلوبة"

        # State how many clause-law results the figures rest on, so a score
        # driven by one applicable clause is not read as a verdict on the
        # whole contract.
        if lang == 'ar':
            basis = f" (بناءً على {n_applicable} من {n_checked} نتيجة بند/نظام منطبقة)" if n_checked else ""
            return f"العقد: {name}\nالحالة: {status_ar}\nالمخاطر: {overall_risk:.0f}%{basis}\nالتوصية: {action_ar}"
        basis = f" (based on {n_applicable} of {n_checked} clause-law results that apply)" if n_checked else ""
        return f"Contract: {name}\nStatus: {status_en}\nRisk: {overall_risk:.0f}%{basis}\nRecommendation: {action_en}"


def _risk_level(risk: Optional[float]) -> Optional[RiskLevel]:
    if risk is None:
        return None
    return (RiskLevel.LOW if risk < 30 else
            RiskLevel.MEDIUM if risk < 50 else
            RiskLevel.HIGH if risk < 75 else RiskLevel.CRITICAL)


# ============================================================
# 5. Main analyzer
# ============================================================

class ContractAnalyzer:
    def __init__(self, embedder: SentenceTransformer):
        self.embedder = embedder
        self.law_kg = LawKnowledgeGraph(embedder)
        self.parser = ContractParser()
        self.engine = ReasoningEngine(embedder)

    def add_law_article(self, article: LawArticle):
        self.law_kg.add_article(article)

    def build_law_index(self):
        self.law_kg.build_index()

    def analyze(self, contract_text: str, contract_name: str, laws: List[str],
                keep_clauses: bool = False):
        # Random per-request id: an id derived from the text could collide
        # for two contracts with the same opening words (and their PDFs
        # would overwrite each other).
        contract_id = f"C_{uuid.uuid4().hex[:8]}"

        # Per-request contract graph -- dropped when this method returns, so
        # nothing from the uploaded contract is kept in server memory.
        contract_kg = ContractKnowledgeGraph(self.embedder)
        contract_kg.add_contract(contract_id, contract_name)

        report_lang = _detect_lang(contract_text)

        parsed = self.parser.parse_text(contract_text)
        clauses = []
        for p in parsed:
            clause = ContractClause(
                id=f"{contract_id}_CL{p['number']}",
                contract_id=contract_id,
                number=p['number'],
                title=p['title'],
                content=p['content'],
                clause_type=p['type'],
                keywords=self.parser.extract_keywords(p['content'])
            )
            contract_kg.add_clause(clause)
            clauses.append(clause)

        law_results: Dict[str, List[AnalysisResult]] = {law: [] for law in laws}

        for clause in clauses:
            query = f"{clause.title} {clause.content}"
            for law_id in laws:
                matches = self.law_kg.search(query, top_k=3, law_filter=[law_id],
                                             query_emb=clause.embedding)
                result = self.engine.analyze(clause, matches, law_id, self.law_kg, lang=report_lang)
                contract_kg.add_match(clause.id, law_id, result)
                law_results[law_id].append(result)

        law_scores: Dict[str, LawScore] = {}
        for law_id in laws:
            results = law_results[law_id]
            if not results:
                continue

            relevant = [r for r in results if r.status != ComplianceStatus.NOT_APPLICABLE]
            compliant = sum(1 for r in relevant if r.status == ComplianceStatus.COMPLIANT)
            partial = sum(1 for r in relevant if r.status == ComplianceStatus.PARTIAL)
            non_comp = sum(1 for r in relevant if r.status == ComplianceStatus.NON_COMPLIANT)
            n_a = len(results) - len(relevant)

            # Compliance and risk are computed over the clauses that relate
            # to this law only; "not applicable" clauses neither earn nor
            # lose credit.
            if relevant:
                compliance_pct = ((compliant + partial * 0.5) / len(relevant)) * 100
                avg_risk = float(np.mean([r.risk_score for r in relevant]))
            else:
                compliance_pct, avg_risk = None, None

            findings = [i for r in relevant for i in r.issues]

            reasoning = (
                f"فحص {len(results)} بند: {compliant}✅ {partial}⚠️ {non_comp}❌ {n_a}➖"
                if report_lang == 'ar' else
                f"Checked {len(results)} clauses: {compliant}✅ {partial}⚠️ {non_comp}❌ {n_a}➖"
            )
            law_scores[law_id] = LawScore(
                law_id=law_id, compliance_pct=compliance_pct, risk_score=avg_risk,
                risk_level=_risk_level(avg_risk), total_checked=len(results),
                compliant=compliant, partial=partial, non_compliant=non_comp,
                not_applicable=n_a, reasoning=reasoning,
                findings=_unique(findings)[:5]
            )

        scored = [lid for lid, s in law_scores.items() if s.compliance_pct is not None]
        if scored:
            total_weight = sum(LAWS[lid]['weight'] for lid in scored)
            overall_compliance = sum(law_scores[lid].compliance_pct * LAWS[lid]['weight'] for lid in scored) / total_weight
            overall_risk = sum(law_scores[lid].risk_score * LAWS[lid]['weight'] for lid in scored) / total_weight
        else:
            overall_compliance, overall_risk = None, None

        n_checked = sum(len(law_results[l]) for l in laws)
        n_applicable = sum(1 for l in laws for r in law_results[l]
                           if r.status != ComplianceStatus.NOT_APPLICABLE)
        summary = self.engine.generate_summary(overall_risk, contract_name, lang=report_lang,
                                               n_applicable=n_applicable, n_checked=n_checked)

        # Ordered clause by clause (then law), so the sampled list covers the
        # whole contract instead of being filled by the first law alone.
        all_risks = [i for ci in range(len(clauses)) for law_id in laws
                     for i in law_results[law_id][ci].issues]
        all_recs = [x for ci in range(len(clauses)) for law_id in laws
                    for x in law_results[law_id][ci].recommendations]

        report = FullReport(
            contract_id=contract_id, contract_name=contract_name,
            timestamp=datetime.now().strftime("%Y-%m-%d %H:%M"), total_clauses=len(clauses),
            laws_checked=laws, law_scores=law_scores,
            overall_compliance=overall_compliance, overall_risk=overall_risk,
            overall_level=_risk_level(overall_risk), summary=summary,
            risks=_unique(all_risks)[:10], recommendations=_unique(all_recs)[:10],
            learned_keywords=_unique(
                kw for kws in contract_kg.learned_patterns.values() for kw in kws
            )[:20],
            trace=contract_kg.trace(contract_id, laws),
            language=report_lang
        )
        if keep_clauses:
            return report, clauses
        return report


# ============================================================
# 6. Exploratory ML layer (kept disabled by default; see thesis Ch. 4)
# ============================================================

class TrainingDataManager:
    def __init__(self, test_size: float = 0.15, val_size: float = 0.15, random_state: int = 42):
        self.test_size = test_size
        self.val_size = val_size
        self.random_state = random_state
        self.label_encoder = LabelEncoder()
        self.scaler = StandardScaler()

    def prepare_data(self, X: np.ndarray, y: List[str], scale: bool = True) -> TrainingData:
        y_encoded = self.label_encoder.fit_transform(y)

        X_temp, X_test, y_temp, y_test = train_test_split(
            X, y_encoded, test_size=self.test_size, random_state=self.random_state,
            stratify=y_encoded if len(set(y_encoded)) > 1 else None
        )
        val_ratio = self.val_size / (1 - self.test_size)
        X_train, X_val, y_train, y_val = train_test_split(
            X_temp, y_temp, test_size=val_ratio, random_state=self.random_state,
            stratify=y_temp if len(set(y_temp)) > 1 else None
        )

        if scale:
            X_train = self.scaler.fit_transform(X_train)
            X_val = self.scaler.transform(X_val)
            X_test = self.scaler.transform(X_test)

        return TrainingData(
            X_train=X_train, X_val=X_val, X_test=X_test,
            y_train=y_train, y_val=y_val, y_test=y_test,
            label_encoder=self.label_encoder, scaler=self.scaler
        )


class ModelTrainer:
    MODELS = {
        "logistic": lambda: LogisticRegression(max_iter=1000, random_state=42),
        "random_forest": lambda: RandomForestClassifier(n_estimators=100, random_state=42),
        "gradient_boosting": lambda: GradientBoostingClassifier(n_estimators=100, random_state=42),
        "svm": lambda: SVC(kernel='rbf', probability=True, random_state=42),
    }

    def __init__(self):
        self.trained_models: Dict[str, any] = {}
        self.results: Dict[str, Dict] = {}

    def train(self, data: TrainingData, model_name: str = "random_forest") -> Dict:
        model = self.MODELS[model_name]()
        model.fit(data.X_train, data.y_train)
        results = {
            "model_name": model_name,
            "train": self._evaluate(model, data.X_train, data.y_train),
            "val": self._evaluate(model, data.X_val, data.y_val),
            "test": self._evaluate(model, data.X_test, data.y_test),
        }
        self.trained_models[model_name] = model
        self.results[model_name] = results
        return results

    def _evaluate(self, model, X, y) -> Dict:
        y_pred = model.predict(X)
        return {
            "accuracy": accuracy_score(y, y_pred),
            "precision": precision_score(y, y_pred, average='weighted', zero_division=0),
            "recall": recall_score(y, y_pred, average='weighted', zero_division=0),
            "f1": f1_score(y, y_pred, average='weighted', zero_division=0),
            "confusion_matrix": confusion_matrix(y, y_pred),
        }

    def predict(self, X: np.ndarray, model_name: str = "random_forest") -> np.ndarray:
        return self.trained_models[model_name].predict(X)


class MLContractAnalyzer:
    """Wraps ContractAnalyzer with an optional, exploratory ML risk classifier.

    Per the thesis (Chapter 4), this classifier is exploratory: it is not
    trained by default in this deployment, and no headline metric is
    reported from it. ``analyze_with_ml`` falls back to the rule-based
    result alone (``ml_prediction=None``) until a model is explicitly
    trained on labeled data.
    """

    def __init__(self, embedder: SentenceTransformer):
        self.analyzer = ContractAnalyzer(embedder)
        self.data_manager = TrainingDataManager()
        self.trainer = ModelTrainer()
        self.is_trained = False

    def analyze_with_ml(self, contract_text: str, contract_name: str, laws: List[str]) -> Tuple[FullReport, Optional[str]]:
        report, clauses = self.analyzer.analyze(contract_text, contract_name, laws, keep_clauses=True)

        ml_prediction = None
        if self.is_trained:
            embeddings = [c.embedding for c in clauses if c.embedding is not None]
            if embeddings:
                X = self.data_manager.scaler.transform(np.array(embeddings))
                predictions = self.trainer.predict(X)
                from collections import Counter
                ml_prediction = Counter(predictions).most_common(1)[0][0]
                ml_prediction = self.data_manager.label_encoder.inverse_transform([ml_prediction])[0]

        return report, ml_prediction


# ============================================================
# 7. PDF report export
# ============================================================

def _fmt_pct(v: Optional[float]) -> str:
    return "—" if v is None else f"{v:.0f}"


def _report_chart(report: FullReport, is_ar: bool, font: str, width: float):
    """Grouped bar chart of compliance % and risk % for each law that has at
    least one applicable clause -- the same figure the web page shows --
    drawn with reportlab's own graphics so no image export library is
    needed. Returns None when no law has an applicable clause."""
    charted = [lid for lid, s in report.law_scores.items() if s.compliance_pct is not None]
    if not charted:
        return None
    if is_ar:
        charted = list(reversed(charted))      # categories read right-to-left
    green = colors.HexColor(STATUS_COLORS["compliance"])
    red = colors.HexColor(STATUS_COLORS["risk"])
    grey = colors.HexColor('#666666')
    txt = (lambda s: _ar(s)) if is_ar else (lambda s: s)

    height = 230
    left, right, bottom, top = 30, 10, 58, 30
    plot_w, plot_h = width - left - right, height - bottom - top
    d = Drawing(width, height)

    # Horizontal grid and y-axis labels, 0 to 100.
    for v in range(0, 101, 20):
        y = bottom + plot_h * v / 100
        d.add(Line(left, y, left + plot_w, y, strokeColor=colors.HexColor('#E3E8E6'),
                   strokeWidth=0.5))
        d.add(String(left - 4, y - 3, str(v), fontName='Helvetica', fontSize=7,
                     fillColor=grey, textAnchor='end'))
    d.add(Line(left, bottom, left + plot_w, bottom, strokeColor=grey, strokeWidth=0.8))

    group_w = plot_w / len(charted)
    bar_w = min(34, group_w * 0.3)
    for i, lid in enumerate(charted):
        s = report.law_scores[lid]
        cx = left + group_w * (i + 0.5)
        pairs = [(s.compliance_pct, green), (s.risk_score, red)]
        if is_ar:
            pairs.reverse()                    # compliance bar on the right
        for j, (val, col) in enumerate(pairs):
            val = max(0.0, min(100.0, float(val or 0)))
            x = cx - bar_w - 1 + j * (bar_w + 2)
            h = plot_h * val / 100
            d.add(Rect(x, bottom, bar_w, h, fillColor=col, strokeColor=None))
            d.add(String(x + bar_w / 2, bottom + h + 3, f"{val:.0f}%", fontName='Helvetica',
                         fontSize=7.5, fillColor=colors.black, textAnchor='middle'))
        # Law name under its group, wrapped onto at most two lines.
        name = LAWS.get(lid, {}).get('ar' if is_ar else 'en', lid)
        words, lines, cur = name.split(), [], ''
        for w in words:
            cand = f"{cur} {w}".strip()
            if cur and pdfmetrics.stringWidth(arabic_reshaper.reshape(cand) if is_ar else cand,
                                              font, 7.5) > group_w - 6:
                lines.append(cur)
                cur = w
            else:
                cur = cand
        lines.append(cur)
        if len(lines) > 2:
            lines = [lines[0], " ".join(lines[1:])]
        for k, line in enumerate(lines):
            d.add(String(cx, bottom - 12 - k * 10, txt(line), fontName=font, fontSize=7.5,
                         fillColor=colors.black, textAnchor='middle'))

    # Legend above the plot.
    items = ([("التوافق %", green), ("المخاطر % (الشدة)", red)] if is_ar
             else [("Compliance %", green), ("Risk % (severity)", red)])
    if is_ar:
        # Right-to-left: first item on the right, swatch to the right of its label.
        rx = width / 2 + 90
        for label, col in items:
            d.add(Rect(rx - 9, height - 16, 9, 9, fillColor=col, strokeColor=None))
            d.add(String(rx - 13, height - 15, txt(label), fontName=font, fontSize=8,
                         fillColor=colors.black, textAnchor='end'))
            rx -= 110
    else:
        lx = width / 2 - 90
        for label, col in items:
            d.add(Rect(lx, height - 16, 9, 9, fillColor=col, strokeColor=None))
            d.add(String(lx + 13, height - 15, label, fontName=font, fontSize=8,
                         fillColor=colors.black))
            lx += 110
    return d


def save_pdf_report(report: FullReport, path: str) -> str:
    """Writes the PDF report in a single language -- Arabic throughout if the
    contract was Arabic, English throughout if it was English -- using the
    bundled Amiri font for Arabic. Contents: overall result, per-law table,
    a chart of compliance and risk per law, the clause-by-clause
    traceability table (which article each clause was
    matched to, and the numbers behind each verdict), and the findings."""
    is_ar = (report.language == 'ar')
    if is_ar:
        font = _ARABIC_FONT if _ARABIC_FONT_OK else "Helvetica"
        font_bold = _ARABIC_FONT_BOLD if _ARABIC_FONT_OK else "Helvetica-Bold"
    else:
        font, font_bold = "Helvetica", "Helvetica-Bold"
    accent = colors.HexColor('#0E4A3F')
    align = TA_RIGHT if is_ar else TA_LEFT

    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=36, bottomMargin=36,
                             leftMargin=36, rightMargin=36)
    styles = getSampleStyleSheet()
    story = []

    page_w = A4[0] - 72

    def P(text, style, width: float = page_w) -> Paragraph:
        """Builds a Paragraph safely in either language.

        Text is escaped for reportlab's markup parser (an '&' or '<' in a
        contract name or clause previously broke PDF generation). Arabic is
        wrapped into lines *before* bidi reordering -- reportlab wraps the
        already-reordered string, which printed multi-line Arabic cells
        with their lines in reverse order."""
        text = str(text)
        if not is_ar:
            return Paragraph(_xml_escape(text), style)
        avail = max(width - 14, 20)
        lines, cur = [], ''
        for w in text.split():
            cand = f"{cur} {w}".strip()
            if cur and pdfmetrics.stringWidth(arabic_reshaper.reshape(cand),
                                              style.fontName, style.fontSize) > avail:
                lines.append(cur)
                cur = w
            else:
                cur = cand
        if cur:
            lines.append(cur)
        return Paragraph('<br/>'.join(_xml_escape(_ar(l)) for l in lines), style)

    def rtl(row):
        """Arabic tables read right-to-left: first column on the right."""
        return list(reversed(row)) if is_ar else row

    title_style = ParagraphStyle('Title', parent=styles['Title'], fontName=font_bold,
                                  alignment=TA_CENTER, fontSize=17, textColor=accent,
                                  spaceAfter=2)
    meta_style = ParagraphStyle('Meta', parent=styles['Normal'], fontName=font,
                                 alignment=align, fontSize=9.5,
                                 leading=16 if is_ar else 14, textColor=colors.HexColor('#333333'))
    stat_style = ParagraphStyle('Stat', parent=styles['Heading2'],
                                 fontName=font_bold, alignment=align, fontSize=12, textColor=colors.black)
    h3_style = ParagraphStyle('H3', parent=styles['Heading3'], fontName=font_bold,
                               alignment=align, fontSize=11, textColor=accent,
                               spaceBefore=4)
    bullet_style = ParagraphStyle('Bullet', parent=styles['Normal'], fontName=font,
                                   alignment=align, fontSize=9, leading=15 if is_ar else 12,
                                   rightIndent=10 if is_ar else 0, leftIndent=0 if is_ar else 10,
                                   spaceBefore=2, textColor=colors.HexColor('#444444'))
    footer_style = ParagraphStyle('Footer', parent=styles['Normal'],
                                   fontName=font if is_ar else 'Helvetica-Oblique',
                                   alignment=align, fontSize=7.5, leading=12 if is_ar else 10,
                                   textColor=colors.grey)
    warn_style = ParagraphStyle('Warn', parent=meta_style, textColor=colors.HexColor('#8A5A00'))
    cell_style = ParagraphStyle('Cell', parent=meta_style, fontSize=8, leading=12 if is_ar else 10)
    cell_c = ParagraphStyle('CellC', parent=cell_style, alignment=TA_CENTER)
    header_style = ParagraphStyle('Header', parent=styles['Normal'], fontSize=8, leading=11,
                                   alignment=TA_CENTER, textColor=colors.white, fontName=font_bold)

    def table_style():
        return TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), accent),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f6fb')]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ])

    def divider():
        story.append(Spacer(1, 6))
        story.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor('#dddddd')))
        story.append(Spacer(1, 6))

    # --- Header -----------------------------------------------------
    story.append(P("تقرير فحص امتثال العقد" if is_ar else "Contract Compliance Report", title_style))
    story.append(Spacer(1, 12))
    story.append(P(f"العقد: {report.contract_name}" if is_ar
                   else f"Contract: {report.contract_name}", meta_style))
    story.append(Spacer(1, 3))
    story.append(P(f"تاريخ الإنشاء: {report.timestamp}" if is_ar
                   else f"Generated: {report.timestamp}", meta_style))
    story.append(Spacer(1, 3))
    story.append(P(f"عدد البنود المحللة: {report.total_clauses}" if is_ar
                   else f"Clauses analyzed: {report.total_clauses}", meta_style))
    if report.contract_type:
        story.append(Spacer(1, 3))
        story.append(P(f"نوع العقد (كما حدده المستخدم): {report.contract_type}" if is_ar
                       else f"Contract type (as stated by the user): {report.contract_type}", meta_style))
    for n in report.scope_notes:
        story.append(Spacer(1, 3))
        story.append(P(("تنبيه: " if is_ar else "Note: ") + n, warn_style))
    divider()

    n_all = len(report.trace)
    n_app = sum(1 for r in report.trace if r.status != ComplianceStatus.NOT_APPLICABLE)
    if report.overall_compliance is None:
        overall = ("لم يُعثر على بنود تتعلق بالأنظمة المختارة" if is_ar
                   else "No clause was found relevant to the selected laws")
    elif is_ar:
        overall = (f"الامتثال الإجمالي: {report.overall_compliance:.0f}%   "
                   f"المخاطر الإجمالية: {report.overall_risk:.0f}% ({report.overall_level.arabic})   "
                   f"البنود المنطبقة: {n_app} من {n_all}")
    else:
        overall = (f"Overall compliance: {report.overall_compliance:.0f}%   "
                   f"Overall risk: {report.overall_risk:.0f}% ({report.overall_level.name.title()})   "
                   f"Applicable results: {n_app} of {n_all}")
    story.append(P(overall, stat_style))
    story.append(Spacer(1, 10))

    # --- Per-law table ------------------------------------------------
    head = (["النظام", "التوافق %", "المخاطر %", "المستوى", "متوافق / جزئي / غير متوافق", "غير منطبق"]
            if is_ar else
            ["Law", "Compliance %", "Risk %", "Level", "Compl. / Partial / Non-compl.", "Not applic."])
    widths = [150, 62, 52, 62, 125, 72]
    rows = [rtl([P(h, header_style, w) for h, w in zip(head, widths)])]
    for lid, s in report.law_scores.items():
        law_name = LAWS.get(lid, {}).get('ar' if is_ar else 'en', lid)
        if s.risk_level is None:
            level_name = "—"
        else:
            level_name = s.risk_level.arabic if is_ar else s.risk_level.value.title()
        # Counts are listed in the same reading order as the header
        # (right-to-left in Arabic, so the compliant count sits on the right).
        counts = (f"{s.non_compliant} / {s.partial} / {s.compliant}" if is_ar
                  else f"{s.compliant} / {s.partial} / {s.non_compliant}")
        rows.append(rtl([
            P(law_name, cell_style, widths[0]),
            P(_fmt_pct(s.compliance_pct), cell_c),
            P(_fmt_pct(s.risk_score), cell_c),
            P(level_name, cell_c, widths[3]),
            P(counts, cell_c),
            P(str(s.not_applicable), cell_c),
        ]))
    tbl = Table(rows, hAlign='CENTER', colWidths=rtl(widths), repeatRows=1)
    tbl.setStyle(table_style())
    story.append(tbl)
    story.append(Spacer(1, 12))

    # --- Chart: compliance and risk per law -----------------------------
    chart = _report_chart(report, is_ar, font, page_w)
    if chart is not None:
        story.append(P("التوافق والمخاطر لكل نظام" if is_ar
                       else "Compliance and risk per law", h3_style))
        story.append(Spacer(1, 4))
        story.append(chart)
        story.append(Spacer(1, 10))

    # --- Clause traceability table -------------------------------------
    if report.trace:
        story.append(P("تتبع البنود: المادة المطابقة وأساس الحكم" if is_ar
                       else "Clause traceability: matched article and basis of each verdict", h3_style))
        story.append(Spacer(1, 4))
        head = (["البند", "النظام", "المادة المطابقة", "التشابه %", "المصطلحات", "المجمع %", "الحكم", "مصطلحات غير مذكورة"]
                if is_ar else
                ["Clause", "Law", "Matched article", "Sim. %", "Terms", "Comb. %", "Verdict", "Terms not mentioned"])
        widths = ([82, 70, 92, 38, 48, 42, 58, 93] if is_ar
                  else [80, 70, 90, 36, 40, 40, 68, 99])
        rows = [rtl([P(h, header_style, w) for h, w in zip(head, widths)])]
        lang = 'ar' if is_ar else 'en'
        for r in report.trace:
            law_short = LAWS.get(r.law_id, {}).get(lang, r.law_id)
            terms = (f"{r.matched_terms} من {r.total_terms}" if is_ar
                     else f"{r.matched_terms}/{r.total_terms}")
            art = (f"م {r.article_number}: {r.article_title}" if is_ar
                   else f"Art. {r.article_number}: {r.article_title}")
            missing = ", ".join(r.missing_terms[:3]) if r.status != ComplianceStatus.COMPLIANT else ""
            rows.append(rtl([
                P(r.clause_title[:60], cell_style, widths[0]),
                P(law_short, cell_style, widths[1]),
                P(art, cell_style, widths[2]),
                P(f"{r.best_sim*100:.0f}", cell_c),
                P(terms, cell_c, widths[4]),
                P(f"{r.combined*100:.0f}", cell_c),
                P(r.status.label(lang), cell_c, widths[6]),
                P(missing or "—", cell_style, widths[7]),
            ]))
        tbl = Table(rows, hAlign='CENTER', colWidths=rtl(widths), repeatRows=1)
        tbl.setStyle(table_style())
        story.append(tbl)
        story.append(Spacer(1, 10))

    def item_list(heading: str, items):
        story.append(P(heading, h3_style))
        story.append(Spacer(1, 4))
        for item in items:
            story.append(P(f"• {item}", bullet_style, page_w - 10))
        story.append(Spacer(1, 6))

    if report.risks:
        item_list("أبرز الملاحظات (عينة):" if is_ar else "Issues found (sample):", report.risks[:10])
    if report.recommendations:
        item_list("التوصيات (عينة):" if is_ar else "Recommendations (sample):", report.recommendations[:10])

    divider()
    story.append(P(
        (
            "هذه النتيجة إشارة فحص آلية للمراجعة وليست حكمًا قانونيًا. هذا نموذج أولي بحثي "
            "(المرحلة الأولى من خارطة طريق تؤجل دمج نموذج لغوي توليدي لمرحلة لاحقة) أُعدّ لرسالة "
            "ماجستير في نظم المعلومات بجامعة الملك عبدالعزيز."
        ) if is_ar else
        ("This output is an automated screening signal for review, not a legal determination. "
          "Research prototype, Phase 1 (no generative LLM); produced for a master's thesis in "
          "Information Systems at King Abdulaziz University."),
        footer_style
    ))

    doc.build(story)
    return path



# ============================================================
# 8. Bootstrap: embedder, analyzer, sample law content
# ============================================================

print("Loading sentence-transformer model (this runs once at startup)...")
_embedder = SentenceTransformer(EMBED_MODEL_NAME)
ml_analyzer = MLContractAnalyzer(_embedder)

# Law content: illustrative articles verified one by one (Oct 2026) against
# the official texts -- the Arabic text published by the Bureau of Experts /
# competent ministry (Arabic is the only binding version) and the Bureau of
# Experts' official English translation -- in their current, amended form
# (e.g. Labor Law as amended by Royal Decree M/44 of 8/2/1446H). Every
# article appears in both languages under its official number, so Arabic and
# English contracts are screened against the same set. Content is a close,
# condensed restatement of each article (not the full text); required terms
# are taken from the article's own official wording so that they match the
# vocabulary used in contracts. Arabic required terms are matched after
# normalization (see normalize_text), so hamza/ta marbuta variants match.
#
# Sources: Personal Data Protection Law (M/19, 9/2/1443H, as amended in
# 1444H); Anti-Cyber Crime Law (M/17, 8/3/1428H); E-Commerce Law
# (M/126, 7/11/1440H); Labor Law (M/51, 23/8/1426H, as amended up to M/44,
# 8/2/1446H); Companies Law (M/132, 1/12/1443H); Anti-Commercial Fraud Law
# (M/19, 23/4/1429H).
SAMPLE_LAWS = {
    "PDPL": [
        ("4", "حقوق صاحب البيانات الشخصية", "لصاحب البيانات الشخصية الحق في العلم بالمسوغ النظامي لجمع بياناته والغرض منه، والحق في الوصول إلى بياناته لدى جهة التحكم والحصول عليها بصيغة مقروءة، وطلب تصحيحها أو إتمامها أو تحديثها، وطلب إتلافها متى انتهت الحاجة إليها.", ["صاحب البيانات", "الوصول", "تصحيح", "تحديث", "إتلاف"]),
        ("4", "Data Subject Rights", "The data subject has the right to be informed of the legal basis and purpose of collecting their personal data, to access it and obtain it in a readable format, to request its correction, completion or updating, and to request its destruction when no longer needed.", ["data subject", "access", "correction|correct", "update|updat", "destruction|destroy|destr"]),
        ("5", "الموافقة على المعالجة", "لا تجوز معالجة البيانات الشخصية أو تغيير الغرض من معالجتها إلا بعد موافقة صاحبها، فيما عدا الأحوال المنصوص عليها في النظام، ويجوز لصاحب البيانات الرجوع عن موافقته في أي وقت.", ["موافقة", "معالجة", "الغرض", "الرجوع عن الموافقة"]),
        ("5", "Consent to Processing", "Personal data may not be processed, nor the purpose of processing changed, without the data subject's consent, except in the cases stated in the Law; the data subject may withdraw consent at any time.", ["consent", "processing|process", "purpose", "withdraw|withdrawal"]),
        ("8", "اختيار جهة المعالجة", "على جهة التحكم ألا تختار إلا جهة معالجة توفر الضمانات اللازمة لتنفيذ أحكام النظام واللوائح، وأن تتحقق من التزامها بها، دون إخلال بمسؤوليتها تجاه صاحب البيانات، بما في ذلك أحكام أي تعاقدات لاحقة تقوم بها جهة المعالجة.", ["جهة المعالجة", "ضمانات", "التحقق", "التزام", "تعاقد"]),
        ("8", "Selection of Processors", "The controller shall only select processors providing the necessary guarantees to implement the Law and its Regulations, and shall monitor their compliance, without prejudice to its own responsibilities, including for any subsequent contracts made by the processor.", ["processor", "guarantees|guarantee", "monitor", "compliance|compliant|comply", "subcontract|sub-contract|subsequent contract"]),
        ("11", "ضوابط جمع البيانات", "يجب أن يكون الغرض من جمع البيانات الشخصية ذا علاقة مباشرة بأغراض جهة التحكم، وأن تكون طرق جمعها واضحة وآمنة وخالية من الخداع، وأن يقتصر محتواها على الحد الأدنى اللازم لتحقيق الغرض، وأن يتوقف جمعها وتتلف متى لم تعد ضرورية.", ["الغرض", "الحد الأدنى", "جمع", "إتلاف"]),
        ("11", "Collection Limitation", "The purpose of collecting personal data shall be directly related to the controller's purposes; collection methods shall be clear, secure and free of deception; the data shall be limited to the minimum necessary for the purpose; and collection shall stop and the data be destroyed once no longer necessary.", ["purpose", "minimum", "collection|collect", "destruction|destroy|destr"]),
        ("18", "إتلاف البيانات والاحتفاظ بها", "على جهة التحكم إتلاف البيانات الشخصية دون تأخير بعد انتهاء الغرض من جمعها، ما لم تُزل منها ما يؤدي إلى معرفة صاحبها، أو يوجد مسوغ نظامي يوجب الاحتفاظ بها مدة محددة، أو تكن متصلة بقضية منظورة.", ["إتلاف", "انتهاء الغرض", "الاحتفاظ", "مدة"]),
        ("18", "Destruction and Retention", "The controller shall destroy personal data without undue delay once the purpose of collection ends, unless identifying elements are removed, a legal basis requires retention for a specified period, or the data relates to a pending judicial case.", ["destruction|destroy|destr", "retention period|retain for", "retain|retention", "period"]),
        ("19", "حماية البيانات", "على جهة التحكم اتخاذ ما يلزم من إجراءات ووسائل تنظيمية وإدارية وتقنية تضمن المحافظة على البيانات الشخصية، بما في ذلك عند نقلها.", ["إجراءات", "تقنية", "المحافظة على البيانات", "حماية"]),
        ("19", "Security Measures", "The controller shall implement the necessary organizational, administrative and technical measures to protect personal data, including during its transfer.", ["measures", "technical", "protection|protect", "security|secure"]),
        ("20", "الإشعار بتسرب البيانات", "تشعر جهة التحكم الجهة المختصة عند علمها بحدوث تسرب للبيانات الشخصية أو تلفها أو وصول غير مشروع إليها، وتشعر صاحب البيانات إذا كان من شأن ذلك أن يرتب ضررًا على بياناته أو يتعارض مع حقوقه.", ["تسرب", "وصول غير مشروع", "إشعار", "تلف"]),
        ("20", "Breach Notification", "The controller shall notify the competent authority upon knowing of any breach, damage or illegal access to personal data, and shall notify the data subject where this would cause damage to their data or prejudice their rights.", ["breach", "notify|notification", "illegal access", "damage"]),
        ("29", "نقل البيانات خارج المملكة", "يجوز نقل البيانات الشخصية إلى خارج المملكة لأغراض محددة، بشرط ألا يمس ذلك الأمن الوطني، وأن يتوافر في الخارج مستوى مناسب لحماية البيانات، وأن يقتصر النقل على الحد الأدنى من البيانات الذي تدعو إليه الحاجة.", ["نقل", "خارج المملكة", "مستوى مناسب", "الحد الأدنى"]),
        ("29", "Cross-border Transfer", "Personal data may be transferred outside the Kingdom only for specified purposes, provided this does not prejudice national security, an adequate level of protection exists abroad, and the transfer is limited to the minimum data needed.", ["transfer", "outside the kingdom", "adequate", "minimum"]),
    ],
    "CYBER": [
        ("2", "أهداف النظام", "يهدف النظام إلى الحد من الجرائم المعلوماتية بما يحقق الأمن المعلوماتي، وحفظ الحقوق المترتبة على الاستخدام المشروع للحاسبات الآلية والشبكات المعلوماتية، وحماية المصلحة العامة والاقتصاد الوطني.", ["الأمن المعلوماتي", "الاستخدام المشروع", "الشبكة المعلوماتية", "حماية"]),
        ("2", "Objectives", "The Law aims to combat cybercrimes so as to enhance information security, protect rights relating to the legitimate use of computers and information networks, and protect the public interest and the national economy.", ["information security", "legitimate use", "network", "protection|protect"]),
        ("3", "التنصت والدخول غير المشروع", "يعاقب بالسجن مدة لا تزيد على سنة وبغرامة لا تزيد على خمسمائة ألف ريال من يتنصت على ما يرسل عبر الشبكة المعلوماتية أو يلتقطه أو يعترضه دون مسوغ نظامي، أو يدخل دخولًا غير مشروع إلى موقع إلكتروني لتغييره أو إتلافه، أو يمس الحياة الخاصة.", ["تنصت", "اعتراض", "دخول غير مشروع", "موقع إلكتروني", "الحياة الخاصة"]),
        ("3", "Interception and Unauthorized Access", "Imprisonment up to one year and a fine up to 500,000 riyals for intercepting data transmitted through an information network without legal authorization, unauthorized access to a website to change or destroy it, or violation of privacy.", ["interception|intercept", "unauthorized access", "website", "privacy"]),
        ("5", "إتلاف البيانات وتعطيل الأنظمة", "يعاقب بالسجن مدة لا تزيد على أربع سنوات وبغرامة لا تزيد على ثلاثة ملايين ريال من يدخل دخولًا غير مشروع لحذف بيانات خاصة أو تدميرها أو تسريبها أو إتلافها أو تغييرها، أو يوقف الشبكة المعلوماتية أو يعطلها، أو يعيق الوصول إلى الخدمة.", ["بيانات خاصة", "تسريب", "حذف", "تعطيل", "دخول غير مشروع"]),
        ("5", "Damage to Data and Systems", "Imprisonment up to four years and a fine up to three million riyals for unauthorized access to delete, destroy, leak, damage or alter private data, causing an information network to halt, or obstructing access to services.", ["private data", "leak", "deletion|delete", "disruption|disrupt|halt", "unauthorized access"]),
        ("13", "المصادرة والإغلاق", "مع عدم الإخلال بحقوق حسني النية، يجوز الحكم بمصادرة الأجهزة والبرامج والوسائل المستخدمة في الجريمة أو الأموال المحصلة منها، وبإغلاق الموقع الإلكتروني أو مكان تقديم الخدمة إذا كان مصدرًا للجريمة بعلم مالكه.", ["مصادرة", "إغلاق", "الأجهزة", "حسني النية"]),
        ("13", "Confiscation and Closure", "Without prejudice to bona fide parties, the devices, programs and means used in the crime or its proceeds may be confiscated, and the website or service location closed if it was the source of the crime with the owner's knowledge.", ["confiscation|confiscate", "closure", "devices", "bona fide"]),
    ],
    "ECC": [
        ("5", "بيانات المستهلك الشخصية", "لا يجوز لموفر الخدمة الاحتفاظ ببيانات المستهلك الشخصية أو اتصالاته الإلكترونية إلا للمدة التي تقتضيها طبيعة التعامل، ويجب حمايتها والحفاظ على خصوصيتها، ولا يجوز استعمالها لأغراض غير مصرح بها أو الإفصاح عنها لجهة أخرى إلا بموافقة المستهلك.", ["بيانات المستهلك", "الاحتفاظ", "خصوصية", "موافقة", "الإفصاح"]),
        ("5", "Consumer Personal Data", "A service provider shall not retain a consumer's personal data or electronic communications beyond the period required by the transaction, shall protect their confidentiality, and shall not use them for unauthorized purposes or disclose them to others without the consumer's consent.", ["personal data", "retain|retention", "confidential|confidentiality", "consent", "disclose|disclosure"]),
        ("6", "الإفصاح عن موفر الخدمة", "على موفر الخدمة الإفصاح في محله الإلكتروني عن اسمه أو أي بيان مميز له وعنوانه، ووسائل الاتصال به، واسم السجل المقيد فيه ورقمه.", ["الإفصاح", "اسم", "عنوان", "وسائل الاتصال", "السجل التجاري"]),
        ("6", "Service Provider Disclosure", "A service provider shall disclose in his e-shop his name or identifying characteristic and address, his contact information, and the name and number of his commercial registration.", ["disclose|disclosure", "name", "address", "contact", "commercial registration|commercial regist"]),
        ("7", "البيان السابق للتعاقد", "يلتزم موفر الخدمة بتقديم بيان للمستهلك يوضح أحكام العقد وشروطه، يشمل إجراءات إبرام العقد، والخصائص الأساسية للمنتجات أو الخدمات، وإجمالي السعر شاملًا الرسوم والضرائب، وترتيبات الدفع والتسليم والتنفيذ، وبيانات الضمان.", ["أحكام العقد", "السعر", "الدفع", "التسليم", "الضمان"]),
        ("7", "Pre-contract Statement", "A service provider shall give the consumer a statement of the contract's terms and conditions, including the steps to conclude it, the basic characteristics of the products or services, the total price including charges and taxes, payment and delivery arrangements, and warranty information.", ["terms and conditions", "price", "payment", "delivery", "warranty"]),
        ("8", "الفاتورة", "على موفر الخدمة تقديم فاتورة إلى المستهلك بعد إبرام العقد مبين فيها تكاليف كل منتج أو خدمة، وإجمالي السعر شاملًا الرسوم والضرائب، وتاريخ التسليم ومكانه.", ["فاتورة", "إجمالي السعر", "الضرائب", "تاريخ التسليم"]),
        ("8", "Invoice", "Upon conclusion of the contract, the service provider shall give the consumer an itemized invoice stating the total price including charges and taxes, and the date and place of delivery.", ["invoice", "total price", "tax", "date of delivery"]),
        ("13", "فسخ العقد خلال سبعة أيام", "للمستهلك فسخ العقد خلال الأيام السبعة التالية لتاريخ تسلمه المنتج أو لتاريخ التعاقد على الخدمة، ما دام لم يستخدم المنتج أو يستفد من الخدمة، مع استثناءات منها المنتجات المصنعة بناءً على طلبه.", ["فسخ العقد", "سبعة|7 ايام", "تسلم", "لم يستخدم"]),
        ("13", "Right to Rescind", "A consumer may rescind the contract within seven days of receiving the product or of the service contract date, provided he has not used the product or benefited from the service, subject to exceptions such as customized products.", ["rescind", "seven days|7 days|seven (7) days", "receipt", "not used"]),
        ("14", "التأخر في التسليم", "ما لم يتفق على مدة أخرى، للمستهلك فسخ العقد إذا تأخر موفر الخدمة عن التسليم أو التنفيذ مدة تزيد على خمسة عشر يومًا، واسترداد ما دفعه، ما لم يكن التأخير بسبب قوة قاهرة، وعلى موفر الخدمة إبلاغ المستهلك بأي تأخير متوقع.", ["تأخر", "التسليم", "خمسة عشر|15 يوم", "استرداد", "قوة قاهرة"]),
        ("14", "Delay in Delivery", "Unless another period is agreed, the consumer may rescind the contract and recover amounts paid if delivery or performance is delayed more than 15 days, unless due to force majeure; the provider shall notify the consumer of any anticipated delay.", ["delay", "delivery", "15 days|fifteen (15) days|fifteen days", "refund|reimburs", "force majeure"]),
    ],
    "LABOR": [
        ("3", "المساواة في حق العمل", "العمل حق للمواطن، والمواطنون متساوون في حق العمل دون أي تمييز على أساس الجنس أو الإعاقة أو السن أو أي شكل آخر من أشكال التمييز، أثناء أداء العمل أو عند التوظيف أو الإعلان عنه.", ["تمييز", "متساوون", "حق العمل", "التوظيف"]),
        ("3", "Equality in the Right to Work", "Work is the right of every citizen, and all citizens are equal in the right to work without discrimination on the basis of gender, disability, age or any other form, during work, at employment or in advertising.", ["discrimination|discriminat", "equal", "right to work", "employment"]),
        ("51", "كتابة عقد العمل", "يجب أن يكتب عقد العمل من نسختين يحتفظ كل طرف بنسخة، وأن يوثق وفق الأحكام النظامية، ويعد العقد قائمًا ولو كان غير مكتوب، ولكل من الطرفين أن يطلب كتابته في أي وقت.", ["عقد العمل", "نسختين", "مكتوب", "توثيق"]),
        ("51", "Written Employment Contract", "The employment contract shall be executed in duplicate, one copy for each party, and certified as required by law; a contract exists even if not written, and either party may demand at any time that it be in writing.", ["employment contract", "duplicate", "writing", "certified|certif|attest"]),
        ("52", "مضمون عقد العمل", "يكون عقد العمل وفق النموذج الموحد الذي تضعه الوزارة، ويحتوي على اسم صاحب العمل ومكانه، واسم العامل وجنسيته وإثبات شخصيته وعنوان إقامته، والأجر المتفق عليه بما فيه المزايا والبدلات، ونوع العمل ومكانه، وتاريخ الالتحاق، ومدة العقد إن كان محددًا، وحقوق كل طرف والتزاماته.", ["الأجر", "البدلات", "نوع العمل", "مدة العقد", "حقوق"]),
        ("52", "Content of the Employment Contract", "The employment contract shall follow the Ministry's model contract and include the employer's name and place, the worker's name, nationality, identity and residence, the agreed wage including benefits and allowances, type and place of work, start date, term if fixed, and each party's rights and obligations.", ["wage", "allowance", "type of work", "term", "obligations|obligation"]),
        ("53", "فترة التجربة", "إذا كان العامل خاضعًا للتجربة وجب النص على ذلك صراحة في عقد العمل وتحديد مدتها بوضوح، على ألا يزيد مجموعها في جميع الأحوال على مائة وثمانين يومًا، ولكل من الطرفين إنهاء العقد خلالها.", ["تجربة", "مائة وثمانين|مائة وثمانون|180", "صراحة", "إنهاء"]),
        ("53", "Probation Period", "If the worker is subject to probation, the employment contract shall expressly provide for it, and the total probation period shall not exceed 180 days in all cases; either party may terminate the contract within it.", ["probation", "180|one hundred and eighty", "expressly", "termination|terminate"]),
        ("55", "انتهاء العقد محدد المدة وتجديده", "ينتهي العقد محدد المدة بانقضاء مدته، فإن استمر الطرفان في تنفيذه عد مجددًا لمدة غير محددة، وإذا تجدد ثلاث مرات متتالية أو بلغت مدته مع التجديد أربع سنوات أيهما أقل تحول إلى عقد غير محدد المدة.", ["محدد المدة", "تجديد", "انقضاء", "غير محدد المدة"]),
        ("55", "Expiry and Renewal of Fixed-term Contracts", "A fixed-term contract ends when its term expires; if the parties continue to perform it, it is deemed renewed for an indefinite term, and if renewed three consecutive times or the total term reaches four years, whichever is less, it becomes an indefinite-term contract.", ["fixed-term", "renewal|renew", "expiry|expir", "indefinite"]),
        ("75", "الإشعار بإنهاء العقد غير محدد المدة", "إذا كان العقد غير محدد المدة والأجر يدفع شهريًا، جاز لأي من طرفيه إنهاؤه لسبب مشروع بإشعار كتابي قبل ثلاثين يومًا على الأقل إن كان الإنهاء من العامل، وقبل ستين يومًا على الأقل إن كان من صاحب العمل؛ وقبل ثلاثين يومًا لغير الأجر الشهري.", ["إشعار", "كتابي", "سبب مشروع", "ستين|ستون", "ثلاثين|ثلاثون"]),
        ("75", "Notice of Termination", "For an indefinite-term contract with a monthly wage, either party may terminate it for a valid reason by written notice of at least 30 days if by the worker and at least 60 days if by the employer; for non-monthly wages the notice is at least 30 days.", ["notice", "written", "valid reason", "60 days|sixty (60) days|sixty days", "30 days|thirty (30) days|thirty days"]),
        ("76", "التعويض عن مهلة الإشعار", "إذا لم يراع الطرف الذي أنهى العقد غير محدد المدة مهلة الإشعار، التزم بأن يدفع للطرف الآخر مبلغًا مساويًا لأجر العامل عن مهلة الإشعار، ما لم يتفق الطرفان على أكثر من ذلك.", ["مهلة الإشعار", "تعويض", "أجر", "إنهاء"]),
        ("76", "Compensation in Lieu of Notice", "If the party terminating an indefinite-term contract does not observe the notice period, it shall pay the other party an amount equal to the worker's wage for the notice period, unless the parties agree on a greater amount.", ["notice period", "compensation|compensat", "wage", "termination|terminate"]),
        ("77", "التعويض عن الإنهاء لسبب غير مشروع", "ما لم يتضمن العقد تعويضًا محددًا، يستحق المتضرر من إنهاء العقد لسبب غير مشروع أجر خمسة عشر يومًا عن كل سنة خدمة في العقد غير محدد المدة، وأجر المدة الباقية في العقد محدد المدة، على ألا يقل التعويض عن أجر شهرين.", ["سبب غير مشروع", "تعويض", "خمسة عشر يوما|خمسة عشر يوم", "شهرين|شهران"]),
        ("77", "Compensation for Termination without Valid Reason", "Unless the contract specifies compensation, the party harmed by termination for an invalid reason is entitled to 15 days' wage per year of service for indefinite-term contracts, or the wage for the remaining term for fixed-term contracts, and not less than two months' wage.", ["invalid reason", "compensation|compensat", "15 days|fifteen (15) days|fifteen days", "two months|2 months|two (2) months"]),
        ("84", "مكافأة نهاية الخدمة", "إذا انتهت علاقة العمل وجب على صاحب العمل أن يدفع للعامل مكافأة عن مدة خدمته بأجر نصف شهر عن كل سنة من السنوات الخمس الأولى وأجر شهر عن كل سنة تالية، على أساس الأجر الأخير.", ["مكافأة", "نهاية الخدمة", "الأجر الأخير", "سنوات"]),
        ("84", "End-of-Service Award", "Upon the end of the employment relationship, the employer shall pay the worker an end-of-service award of half a month's wage for each of the first five years and one month's wage for each following year, based on the last wage.", ["end-of-service", "award", "last wage", "years"]),
        ("109", "الإجازة السنوية", "يستحق العامل إجازة سنوية بأجر يدفع مقدمًا لا تقل عن واحد وعشرين يومًا، تزاد إلى ثلاثين يومًا إذا أمضى خمس سنوات متصلة، ولا يجوز النزول عنها أو أخذ بدل نقدي عنها أثناء الخدمة.", ["إجازة سنوية", "واحد وعشرين|واحد وعشرون|21", "ثلاثين|ثلاثون", "بأجر"]),
        ("109", "Annual Leave", "A worker is entitled to prepaid annual leave of not less than 21 days, increased to not less than 30 days after five consecutive years of service; it may not be waived or replaced by cash during service.", ["annual leave", "21 days|twenty-one", "30 days|thirty (30) days|thirty days", "paid leave|with pay|full pay"]),
    ],
    "COMMERCIAL": [
        ("4", "أشكال الشركات", "تتخذ الشركة التي تؤسس وفقًا لأحكام النظام أحد الأشكال الآتية: شركة التضامن، أو شركة التوصية البسيطة، أو شركة المساهمة، أو شركة المساهمة المبسطة، أو الشركة ذات المسؤولية المحدودة.", ["شكل الشركة|شكلها", "شركة", "ذات المسؤولية المحدودة", "مساهمة", "تضامن"]),
        ("4", "Forms of Companies", "A company incorporated under the Law shall take one of these forms: general partnership, limited partnership, joint-stock company, simplified joint-stock company, or limited liability company.", ["legal form|form of", "company", "limited liability", "joint-stock", "partnership"]),
        ("8", "قيد وثائق تأسيس الشركة", "يجب أن يكون عقد تأسيس الشركة أو نظامها الأساس وأي تعديل عليه مكتوبًا وإلا كان باطلًا، ويجب قيده لدى السجل التجاري، ويسأل بالتضامن من تسبب في عدم القيد عن التعويض، ولا يحتج على الغير بها إلا بعد القيد.", ["عقد التأسيس", "مكتوب", "السجل التجاري", "قيد", "بالتضامن"]),
        ("8", "Registration of Incorporation Documents", "A company's articles of incorporation or association and their amendments shall be in writing, otherwise void, and shall be registered with the Commercial Register; those causing failure to register are jointly liable for resulting damage, and the documents may not be invoked against third parties before registration.", ["articles of association|articles of incorporation", "writing", "commercial register", "registration", "jointly liable"]),
        ("9", "اكتساب الشخصية الاعتبارية", "تكتسب الشركة الشخصية الاعتبارية بعد قيدها لدى السجل التجاري، وتنتقل إليها العقود التي أجراها المؤسسون لحسابها، ويسأل من تعامل باسمها شخصيًا وبالتضامن إذا لم تستوف إجراءات التأسيس.", ["الشخصية الاعتبارية", "السجل التجاري", "المؤسسون", "التأسيس"]),
        ("9", "Acquisition of Legal Personality", "A company acquires legal personality upon registration with the Commercial Register; contracts concluded by the incorporators on its behalf transfer to it, and those acting in its name are personally and jointly liable if incorporation is not completed.", ["legal personality", "commercial register", "incorporators|incorporator", "incorporation"]),
        ("12", "البيانات الواجب ذكرها في عقود الشركة", "يجب أن يوضع على العقود والمخالصات وغيرها من الوثائق التي تصدرها الشركة: اسم الشركة وشكلها وعنوان مركزها الرئيس وبريدها الإلكتروني ورقم قيدها لدى السجل التجاري، ورأس مالها ومقدار المدفوع منه.", ["اسم الشركة|شركة", "المركز الرئيس|مركز|مقر", "رقم السجل|سجل تجاري رقم|رقم القيد|رقم", "السجل التجاري", "رأس مال"]),
        ("12", "Information Required in Company Documents", "Contracts, clearances and other documents issued by a company shall state the company's name, form, head office address, e-mail and commercial registration number, and its capital and paid-up amount.", ["company name|company", "head office|headquarter|registered office", "commercial registration|c.r. no|cr no|registration no", "capital", "address"]),
        ("35", "شركة التضامن", "شركة التضامن يؤسسها شخصان أو أكثر يكونون مسؤولين شخصيًا في جميع أموالهم وبالتضامن عن ديون الشركة والتزاماتها، ويكتسب الشريك فيها صفة التاجر.", ["تضامن", "مسؤولين شخصيا|مسؤولية شخصية", "ديون", "الشريك"]),
        ("35", "General Partnership", "A general partnership is incorporated by two or more persons who are jointly and personally liable for the company's debts and liabilities; each partner acquires the capacity of a merchant.", ["general partnership", "jointly", "personally liable", "debts"]),
        ("58", "شركة المساهمة", "شركة المساهمة يؤسسها شخص واحد أو أكثر ويكون رأس مالها مقسمًا إلى أسهم قابلة للتداول، وتكون الشركة وحدها مسؤولة عن ديونها والتزاماتها، وتقتصر مسؤولية المساهم على قيمة الأسهم التي اكتتب فيها.", ["مساهمة", "أسهم", "مسؤولية المساهم", "ديون"]),
        ("58", "Joint-Stock Company", "A joint-stock company is incorporated by one or more persons with capital divided into tradable shares; the company alone is liable for its debts, and each shareholder's liability is limited to the value of subscribed shares.", ["joint-stock", "shares", "shareholder", "liability|liable"]),
        ("61", "بيانات النظام الأساس", "يجب أن يشتمل النظام الأساس لشركة المساهمة على اسم الشركة ومركزها الرئيس وغرضها ورأس مالها وعدد أسهمها وأنواعها وقيمتها الاسمية ومدتها وإدارتها وعدد أعضاء مجلس الإدارة وتاريخ بدء السنة المالية وانتهائها.", ["النظام الأساس", "رأس مال", "الأسهم", "مجلس الإدارة", "السنة المالية"]),
        ("61", "Articles of Association", "The articles of association of a joint-stock company shall include its name, head office, purpose, capital, number, types and nominal value of shares, term, management and number of board members, and the start and end of the fiscal year.", ["articles of association", "capital", "shares", "board", "fiscal year"]),
    ],
    "ANTI_FRAUD": [
        ("1", "المنتج المغشوش", "المنتج المغشوش كل منتج دخل عليه تغيير أو عبث أفقده شيئًا من قيمته في ذاته أو طبيعته أو عناصره أو صفاته أو مصدره أو قدره، وكل منتج غير مطابق للمواصفات القياسية المعتمدة، والمنتج الفاسد غير الصالح للاستعمال.", ["مغشوش", "المواصفات القياسية", "مطابق", "فاسد"]),
        ("1", "Fraudulent Product", "A fraudulent product is any product altered or tampered with so as to lose some of its value in its substance, nature, components, characteristics, origin or quantity, any product not conforming to approved standard specifications, and any defective product no longer fit for use.", ["fraudulent", "standard specifications|standard specification", "conformity|conform", "defective"]),
        ("2", "المخالفات: الخداع والغش", "يعد مخالفًا من خدع أو شرع في الخداع في ذاتية المنتج أو طبيعته أو عناصره أو صفاته الجوهرية أو مصدره أو قدره، أو غش المنتج، أو باع منتجًا مغشوشًا أو عرضه، أو صنع أو باع منتجات مخالفة للمواصفات القياسية المعتمدة.", ["خداع", "غش", "مغشوش", "المواصفات", "مصدر المنتج"]),
        ("2", "Violations: Deception and Fraud", "A violator is anyone who deceives or attempts to deceive as to a product's substance, nature, components, essential characteristics, origin or quantity, commits product fraud, sells or displays a fraudulent product, or makes or sells products not conforming to approved standard specifications.", ["deception|deceiv|mislead", "fraud", "origin", "specifications|specification", "quantity"]),
        ("16", "عقوبة المخالفات", "يعاقب بغرامة لا تزيد على خمسمائة ألف ريال أو بالسجن مدة لا تزيد على سنتين أو بهما معًا كل من ارتكب إحدى المخالفات المنصوص عليها في المادة الثانية.", ["غرامة", "خمسمائة ألف|500", "السجن", "عقوبة"]),
        ("16", "Penalty", "Anyone committing a violation under Article 2 shall be punished by a fine not exceeding 500,000 riyals, imprisonment not exceeding two years, or both.", ["fine", "500,000|500000", "imprisonment", "penalty|penalties"]),
        ("22", "سحب المنتج ورد قيمته", "مع عدم الإخلال بحق المتضرر في التعويض، يلتزم المخالف بسحب المنتج المغشوش وإعادة قيمته إلى المشتري وفق ما تحدده اللائحة.", ["سحب", "إعادة قيمته", "التعويض", "المشتري"]),
        ("22", "Withdrawal and Refund", "Without prejudice to the aggrieved party's right to compensation, the violator shall withdraw the fraudulent product and refund buyers as specified in the Regulations.", ["withdraw|withdrawal", "refund|reimburs", "compensation", "buyer"]),
    ],
}

for law_id, articles in SAMPLE_LAWS.items():
    for i, (num, title, content, keywords) in enumerate(articles):
        ml_analyzer.analyzer.add_law_article(LawArticle(
            id=f"{law_id}_{num}_{i}", law_id=law_id, number=num,
            title=title, content=content, keywords=keywords
        ))
ml_analyzer.analyzer.build_law_index()
print(f"Loaded {len(ml_analyzer.analyzer.law_kg.articles)} sample law articles across {len(LAWS)} laws.")


# ============================================================
# 9. Gradio interface
# ============================================================

# Expert-evaluation questionnaire -- the Google Form the user created herself
# (she built it from the question list drafted earlier, since no Google Forms
# connector is available in this session to create it directly).
FEEDBACK_FORM_URL = "https://docs.google.com/forms/d/e/1FAIpQLSdj4pOs-fiJaAw3at0xevJRc4FiihbwYYcGRtZWbrEK_vJQlg/viewform"

# Short illustrative example contracts so an expert can test the tool in one
# click instead of needing to bring or paste their own contract first.
EXAMPLE_CONTRACTS = {
    "ar_tech_services": {
        "name": "عقد تقديم خدمات تقنية (مثال)",
        "text": (
            "عقد تقديم خدمات تقنية\n\n"
            "البند الأول: الأطراف\n"
            "أبرم هذا العقد بين شركة (الطرف الأول) وشركة (الطرف الثاني) لتقديم خدمات "
            "دعم فني وتطوير أنظمة لمدة سنة واحدة قابلة للتجديد.\n\n"
            "البند الثاني: البيانات والخصوصية\n"
            "يلتزم الطرف الثاني بالحفاظ على سرية أي بيانات شخصية يطّلع عليها أثناء "
            "تنفيذ الخدمة، ولا يجوز نقلها لأي جهة خارجية دون إذن كتابي مسبق.\n\n"
            "البند الثالث: الأجر وشروط الدفع\n"
            "يلتزم الطرف الأول بدفع مبلغ العقد على دفعات شهرية خلال خمسة أيام عمل من "
            "تاريخ استلام الفاتورة.\n\n"
            "البند الرابع: الموظفون\n"
            "يتحمل الطرف الثاني كامل الالتزامات تجاه موظفيه المكلفين بتنفيذ هذا العقد "
            "من أجور ومكافآت نهاية خدمة وفق الأنظمة المعمول بها.\n\n"
            "البند الخامس: الإنهاء\n"
            "يجوز لأي من الطرفين إنهاء العقد بإشعار كتابي مدته ثلاثون يومًا."
        ),
    },
    "en_service_agreement": {
        "name": "Service & Maintenance Agreement (example)",
        "text": (
            "SERVICE AND MAINTENANCE AGREEMENT\n\n"
            "Article 1: Parties and Scope\n"
            "This Agreement is entered into between the Client and the Service Provider "
            "for the provision of IT maintenance and support services for a period of "
            "twelve (12) months, renewable by mutual written consent.\n\n"
            "Article 2: Data Protection\n"
            "The Service Provider shall keep confidential any personal data accessed in "
            "the course of performing the Services and shall not transfer such data to "
            "any third party without the Client's prior written consent.\n\n"
            "Article 3: Fees and Payment\n"
            "The Client shall pay the agreed fees on a monthly basis within fifteen (15) "
            "business days of receiving a valid invoice from the Service Provider.\n\n"
            "Article 4: Personnel\n"
            "The Service Provider remains solely responsible for the wages, end-of-service "
            "benefits, and other statutory entitlements of its own employees assigned to "
            "this Agreement.\n\n"
            "Article 5: Termination\n"
            "Either party may terminate this Agreement by giving thirty (30) days' prior "
            "written notice to the other party."
        ),
    },
}

NOTE_MD_EN = """
**Phase 1 research prototype** for a master's thesis in Information Systems
(King Abdulaziz University) — a knowledge-graph and semantic-matching
framework that screens business-contract clauses against six general Saudi
laws, with fixed rule-based reasoning (no generative LLM in this phase).

**How to evaluate**
1. Paste or upload a contract — your own, or use a "Load example" button below
2. Choose the type of your contract, select the laws that govern it, then click **Analyze**
   (for example, the Labor Law for employment contracts, not for a services
   contract between two companies)
3. Review the summary, the **clause traceability table**, the chart and the PDF report
4. Share your judgment as a domain expert via the feedback link at the bottom
   of this page (correctness, clarity, traceability, usability)

**What the results mean**
| Term | Meaning |
|---|---|
| Matched article | The article of the selected law closest in meaning to the clause (via the knowledge graph and semantic search) |
| Similarity % | Semantic similarity between the clause and that article |
| Terms | How many of the matched articles' required terms appear in the clause |
| Verdict | Combined score = 0.4 × terms + 0.6 × similarity: above 55% compliant, above 35% partial, otherwise non-compliant |
| Not applicable | The clause is unrelated to that law (it does not mention what the law regulates, or its similarity is below 47% with fewer than half the required terms) and is left out of that law's figures |
| Compliance % | (compliant + half of partial) ÷ applicable clauses, per law |
| Risk % | A separate severity score from each clause's verdict, averaged per law — **not** simply 100 − compliance |

**Known limitations — please read before judging "correctness"**
- The law content loaded is a selected set of articles per law (78 in total:
  the same 39 articles in Arabic and in English, each verified against the
  official text in its current amended form), not the complete legal texts, so a well-drafted clause may still not reach "Compliant."
  This is an expected, documented limitation — not necessarily a sign the
  tool misjudged the clause.
- The thresholds above are fixed heuristics, not calibrated on labelled data.
- The tool does not decide which laws govern your contract. A clause is
  checked against a law only if it mentions what that law regulates (for
  example personal data, workers, or consumers), and a note appears if a
  selected law does not usually govern your contract type. These checks rely
  on wording, so a clause may still be related to an article that does not
  legally govern that kind of contract. Please judge correctness with this
  in mind.
- The exploratory ML risk classifier (Chapter 4 of the thesis) is **not
  trained** in this deployment; only the rule-based result is shown.
- Scanned PDFs (images of text) cannot be read; use a text PDF, DOCX or paste the text.
- Results are an automated screening signal, not a legal determination.

**Your data**
Your contract is analyzed in memory for your request only and is not saved
to any database or log. The uploaded file and the generated PDF are
temporary files that are deleted automatically from the server within about
an hour. Please avoid uploading documents containing third-party personal
data you are not authorized to share.
"""

NOTE_MD_AR = """
**نموذج أولي بحثي (المرحلة الأولى)** لرسالة ماجستير في نظم المعلومات (جامعة
الملك عبدالعزيز) — يفحص بنود العقود التجارية مقابل ستة أنظمة سعودية عامة
باستخدام رسم معرفي ومطابقة دلالية وقواعد منطقية ثابتة، دون نموذج لغوي توليدي
في هذه المرحلة.

**خطوات التقييم**
1. ألصق/ي نص العقد أو ارفع/ي ملفًا — عقدك الخاص، أو استخدم/ي زر "تحميل مثال" بالأسفل
2. اختر/ي نوع عقدك، ثم الأنظمة التي تحكمه، ثم اضغط/ي **تحليل** (مثلًا: نظام العمل
   لعقود العمل، لا لعقد خدمات بين شركتين)
3. راجع/ي الملخص و**جدول تتبع البنود** والرسم البياني وتقرير PDF
4. شاركنا تقييمك كخبير/ة عبر رابط التقييم أسفل الصفحة (الصحة، الوضوح، قابلية
   التتبع، سهولة الاستخدام)

**معنى النتائج**
| المصطلح | المعنى |
|---|---|
| المادة المطابقة | أقرب مادة في النظام المختار إلى معنى البند (عبر الرسم المعرفي والبحث الدلالي) |
| التشابه % | درجة التشابه الدلالي بين البند وتلك المادة |
| المصطلحات | عدد المصطلحات المطلوبة في المواد المطابقة التي وردت في البند |
| الحكم | الدرجة المجمعة = 0.4 × المصطلحات + 0.6 × التشابه: أعلى من 55% متوافق، أعلى من 35% متوافق جزئيًا، وغير ذلك غير متوافق |
| غير منطبق | البند لا يتعلق بهذا النظام (لا يتناول موضوعه، أو تشابهه أقل من 47% مع أقل من نصف المصطلحات المطلوبة) ويُستبعد من حساب النظام |
| نسبة التوافق | (المتوافق + نصف الجزئي) ÷ البنود المنطبقة، لكل نظام |
| نسبة المخاطر | درجة مستقلة لشدة الخطر مشتقة من حكم كل بند بمتوسطها لكل نظام — **وليست** مجرد 100 ناقص نسبة التوافق |

**قيود معروفة — رجاءً الاطلاع قبل الحكم على "الصحة"**
- الأنظمة المحمّلة مواد مختارة (78 مادة: 39 مادة نفسها بالعربية والإنجليزية، كل
  منها مُطابَقة على النص الرسمي بصيغته النافذة بعد التعديلات) وليست النصوص الكاملة، لذلك قد لا يصل بند مكتوب بصياغة جيدة إلى
  "متوافق" — وهذا قيد موثّق ومتوقع، وليس بالضرورة خطأ من الأداة.
- العتبات أعلاه قيم استدلالية ثابتة، ولم تُعاير على بيانات موسومة.
- الأداة لا تقرر الأنظمة التي تحكم عقدك. لا يُفحص البند مقابل نظام إلا إذا تناول
  موضوعه (مثل البيانات الشخصية أو العمال أو المستهلك)، ويظهر تنبيه إذا اخترت نظامًا
  لا يحكم عادة نوع عقدك. تعتمد هذه الضوابط على الألفاظ، فقد تربط الأداة بندًا بمادة
  لا تحكم هذا النوع من العقود قانونًا. يُرجى مراعاة ذلك عند الحكم على الصحة.
- مصنّف المخاطر بالتعلم الآلي (الفصل الرابع من الرسالة) **غير مدرّب** في هذا
  النشر؛ المعروض هو النتيجة القاعدية فقط.
- ملفات PDF الممسوحة ضوئيًا (صور) لا يمكن قراءتها؛ استخدم/ي PDF نصيًا أو DOCX أو الصق/ي النص.
- النتائج إشارة فحص آلية فقط، وليست حكمًا قانونيًا.

**خصوصية بياناتك**
يُحلَّل العقد في الذاكرة لطلبك فقط، ولا يُحفظ في أي قاعدة بيانات أو سجل. الملف
المرفوع وتقرير PDF ملفات مؤقتة تُحذف تلقائيًا من الخادم خلال ساعة تقريبًا.
يُرجى تجنّب رفع مستندات تحتوي بيانات شخصية لطرف ثالث لا تملك/ين إذنًا بمشاركتها.
"""


_REPORT_DIR = Path(tempfile.gettempdir()) / "contract_reports"
_REPORT_DIR.mkdir(exist_ok=True)
_REPORT_MAX_AGE_S = 15 * 60


def _cleanup_old_reports():
    """Deletes generated PDFs older than 15 minutes (Gradio serves its own
    copy, so ours is only needed for a moment)."""
    now = time.time()
    for f in _REPORT_DIR.glob("report_*.pdf"):
        try:
            if now - f.stat().st_mtime > _REPORT_MAX_AGE_S:
                f.unlink()
        except OSError:
            pass


# Contract types offered in the interface. The user states the type and
# chooses the laws to check; the type does not change the analysis. It is
# shown in the report and recorded so that results can be read, and the
# expert evaluation analyzed, by type of contract.
CONTRACT_TYPES = [
    ("employment", "عقد عمل", "Employment contract"),
    ("services", "عقد خدمات", "Services contract"),
    ("supply", "عقد توريد", "Supply contract"),
    ("construction", "عقد مقاولات أو إنشاءات", "Construction contract"),
    ("consumer_online", "بيع أو خدمة للمستهلك عبر الإنترنت", "Online sale or service to consumers"),
    ("data_it", "اتفاقية معالجة بيانات أو خدمات تقنية", "Data processing or IT services agreement"),
    ("company", "عقد تأسيس شركة أو شراكة", "Company formation or partnership"),
    ("government", "عقد حكومي", "Government contract"),
    ("other", "أخرى", "Other"),
]
_CONTRACT_TYPE_NAMES = {k: (ar, en) for k, ar, en in CONTRACT_TYPES}

# Scope notes: laws that do not usually govern a contract type. The user
# still chooses the laws and the analysis still runs; the note only explains
# why matches against such a law may not be meaningful. The Personal Data
# Protection Law, the Anti-Cyber Crime Law and the Companies Law (whose
# Art. 12 applies to any contract a company issues) can concern any contract
# type and carry no note.
LAW_SCOPE = {
    "LABOR": ({"employment"},
              "ينظم العلاقة بين العامل وصاحب العمل",
              "governs the relationship between an employer and a worker"),
    "ECC": ({"consumer_online"},
            "ينظم تعاملات التجارة الإلكترونية مع المستهلك",
            "governs e-commerce dealings with consumers"),
    "ANTI_FRAUD": ({"supply", "consumer_online", "construction", "government"},
                   "ينظم الغش في المنتجات والسلع",
                   "governs fraud in products and goods"),
}


def _scope_notes(contract_type: Optional[str], laws: List[str], lang: str) -> List[str]:
    if not contract_type or contract_type == "other":
        return []
    ar_type, en_type = _CONTRACT_TYPE_NAMES.get(contract_type, (contract_type, contract_type))
    notes = []
    for lid in laws:
        if lid not in LAW_SCOPE or contract_type in LAW_SCOPE[lid][0]:
            continue
        _, what_ar, what_en = LAW_SCOPE[lid]
        if lang == 'ar':
            notes.append(f"اخترت {LAWS[lid]['ar']} مع {ar_type}. هذا النظام {what_ar}، "
                         f"ولا يحكم عادة هذا النوع من العقود، لذلك قد لا تكون مطابقاته دقيقة.")
        else:
            notes.append(f"You selected the {LAWS[lid]['en']} for a {en_type.lower()}. This law "
                         f"{what_en} and does not usually govern this type of contract, so its "
                         f"matches may not be meaningful.")
    return notes


def _contract_type_label(key: Optional[str], other: Optional[str], lang: str) -> str:
    if not key:
        return ""
    ar, en = _CONTRACT_TYPE_NAMES.get(key, (key, key))
    label = ar if lang == 'ar' else en
    if key == "other" and other and other.strip():
        label = f"{label}: {other.strip()}"
    return label


def _read_uploaded_file(contract_file) -> Tuple[Optional[str], Optional[str]]:
    """Returns (text, file name). Supports both the str path (modern Gradio)
    and object-with-.name (older)."""
    if contract_file is None:
        return None, None
    path = contract_file if isinstance(contract_file, str) else getattr(contract_file, "name", None)
    if not path:
        return None, None
    name = os.path.basename(path)
    try:
        if path.lower().endswith('.pdf'):
            text, _ = ml_analyzer.analyzer.parser.parse_pdf(path)
        elif path.lower().endswith('.docx'):
            text, _ = ml_analyzer.analyzer.parser.parse_docx(path)
        else:
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        return text, name
    except Exception as e:
        raise RuntimeError(f"❌ Could not read the file / تعذرت قراءة الملف ({name}): {e}")


_TRACE_COLS = {
    'ar': ["البند", "النظام", "المادة المطابقة", "التشابه %", "المصطلحات", "المجمع %",
           "الحكم", "مصطلحات غير مذكورة", "نص البند (مقتطف)"],
    'en': ["Clause", "Law", "Matched article", "Similarity %", "Terms", "Combined %",
           "Verdict", "Terms not mentioned", "Clause text (excerpt)"],
}


def _trace_dataframe(report: FullReport) -> pd.DataFrame:
    lang = 'ar' if report.language == 'ar' else 'en'
    rows = []
    for r in report.trace:
        art = (f"م {r.article_number}: {r.article_title}" if lang == 'ar'
               else f"Art. {r.article_number}: {r.article_title}")
        missing = ", ".join(r.missing_terms) if r.status != ComplianceStatus.COMPLIANT else ""
        rows.append([
            f"{r.clause_number}. {r.clause_title}",
            LAWS.get(r.law_id, {}).get(lang, r.law_id),
            art,
            round(r.best_sim * 100),
            f"{r.matched_terms}/{r.total_terms}",
            round(r.combined * 100),
            f"{r.status.icon} {r.status.label(lang)}",
            missing or "—",
            r.clause_excerpt,
        ])
    return pd.DataFrame(rows, columns=_TRACE_COLS[lang])


def _empty_outputs(message: str):
    """Shows an input or processing problem in the summary panel. Each line
    is set in its own direction, so Arabic and English both read correctly."""
    lines = "".join(f'<p dir="auto">{_xml_escape(l.lstrip("❌ ").strip())}</p>'
                    for l in message.split("\n") if l.strip())
    return f'<div class="sum-msg">{lines}</div>', None, None, None


_LEVEL_COLORS = {"low": "#2E7D5B", "medium": "#B7832F", "high": "#C2622D", "critical": "#B3412F"}


def _summary_html(report: "FullReport", source_note: Optional[str], ml_pred) -> str:
    """Summary panel shown above the chart: overall verdict, the basis of
    the figures, and one row per selected law. Written in the contract's
    language and direction; colours mark the risk level."""
    is_ar = (report.language == 'ar')
    d = 'rtl' if is_ar else 'ltr'
    e = _xml_escape
    T = (lambda ar, en: ar) if is_ar else (lambda ar, en: en)
    lines = report.summary.split("\n")
    status = lines[1].split(":", 1)[-1].strip() if len(lines) > 1 else ""
    action = lines[-1].split(":", 1)[-1].strip() if lines else ""
    n_all = len(report.trace)
    n_app = sum(1 for r in report.trace if r.status != ComplianceStatus.NOT_APPLICABLE)
    level = report.overall_level.value if report.overall_level else None
    color = _LEVEL_COLORS.get(level, "#5C6B65")

    meta = []
    if source_note:
        meta.append(f"{T('الملف', 'File')}: {e(source_note)}")
    meta.append(f"{T('نوع العقد', 'Contract type')}: {e(report.contract_type)}")
    meta.append(f"{T('البنود المحللة', 'Clauses analyzed')}: {report.total_clauses}")

    if report.overall_risk is None:
        head = f'<div class="sum-head"><div class="sum-status" style="color:#5C6B65">{e(status)}</div></div>'
    else:
        lvl = report.overall_level.arabic if is_ar else report.overall_level.value.title()
        head = (f'<div class="sum-head" style="border-color:{color}">'
                f'<div class="sum-status" style="color:{color}">{e(status)}</div>'
                f'<div class="sum-figs">{T("التوافق", "Compliance")} <b>{report.overall_compliance:.0f}%</b>'
                f' &nbsp;|&nbsp; {T("المخاطر", "Risk")} <b>{report.overall_risk:.0f}%</b> ({e(lvl)})</div>'
                f'<div class="sum-basis">{T(f"مبنية على {n_app} من {n_all} نتيجة بند/نظام منطبقة", f"Based on {n_app} of {n_all} clause-law results that apply")}'
                f' &nbsp;|&nbsp; {T("التوصية", "Recommendation")}: {e(action)}</div></div>')

    rows = []
    for lid, sc in report.law_scores.items():
        name = e(LAWS[lid]['ar' if is_ar else 'en'])
        counts = (f'{T("متوافق", "compliant")} {sc.compliant} · {T("جزئي", "partial")} {sc.partial} · '
                  f'{T("غير متوافق", "non-compliant")} {sc.non_compliant} · {T("غير منطبق", "not applicable")} {sc.not_applicable}')
        if sc.compliance_pct is None:
            rows.append(f'<tr><td><span class="dot" style="background:#B8C2BE"></span>{name}</td>'
                        f'<td colspan="2">{T("لا توجد بنود منطبقة", "No applicable clauses")}</td><td class="cnt">{counts}</td></tr>')
        else:
            c = _LEVEL_COLORS.get(sc.risk_level.value, "#5C6B65")
            rows.append(f'<tr><td><span class="dot" style="background:{c}"></span>{name}</td>'
                        f'<td>{sc.compliance_pct:.0f}%</td><td style="color:{c}">{sc.risk_score:.0f}%</td>'
                        f'<td class="cnt">{counts}</td></tr>')
    table = (f'<table class="sum-table"><thead><tr><th>{T("النظام", "Law")}</th><th>{T("التوافق", "Compliance")}</th>'
             f'<th>{T("المخاطر", "Risk")}</th><th>{T("توزيع الأحكام", "Verdicts")}</th></tr></thead>'
             f'<tbody>{"".join(rows)}</tbody></table>')
    scope = "".join(f'<p class="sum-warn">⚠️ {e(n)}</p>' for n in report.scope_notes)
    ml = ""
    if ml_pred:
        ml = f'<p class="sum-note">{T("تنبؤ تعلم آلي (استكشافي)", "ML prediction (exploratory)")}: {e(str(ml_pred))}</p>'
    note = T("تفاصيل كل بند والمادة المطابقة له في جدول تتبع البنود أسفل الصفحة.",
             "Each clause, its matched article and the basis of its verdict are in the clause traceability table below.")
    return (f'<div class="sum" dir="{d}"><p class="sum-meta">{" &nbsp;|&nbsp; ".join(meta)}</p>'
            f'{head}{scope}{table}{ml}<p class="sum-note">{note}</p></div>')


def analyze_ui(contract_text, contract_name, contract_type, other_type, selected_laws, contract_file):
    source_note = None
    if contract_file is not None:
        try:
            text, fname = _read_uploaded_file(contract_file)
        except RuntimeError as e:
            return _empty_outputs(str(e))
        if not text or not text.strip():
            # Previously this silently fell back to whatever was in the text
            # box (possibly the loaded example), so the expert would rate a
            # report that was not about their file.
            return _empty_outputs(
                f"❌ لم يُستخرج أي نص من الملف ({fname}). قد يكون PDF ممسوحًا ضوئيًا (صورة). "
                f"ارفع/ي نسخة نصية (PDF نصي أو DOCX) أو الصق/ي النص في المربع بعد إزالة الملف.\n"
                f"❌ No text could be extracted from the file ({fname}). It may be a scanned "
                f"(image) PDF. Upload a text PDF or DOCX, or remove the file and paste the text."
            )
        contract_text = text
        source_note = fname

    if not contract_text or not contract_text.strip():
        return _empty_outputs("❌ Enter contract text or upload a file (PDF/DOCX/TXT) / أدخل/ي نص العقد أو ارفع/ي ملفًا")
    if not contract_type:
        return _empty_outputs("❌ Select the contract type / اختر/ي نوع العقد")
    if not selected_laws:
        return _empty_outputs("❌ Select at least one law / اختر/ي نظامًا واحدًا على الأقل")

    try:
        report, ml_pred = ml_analyzer.analyze_with_ml(
            contract_text, contract_name or source_note or "Contract / عقد", selected_laws)
    except Exception as e:
        print(f"Analysis failed: {type(e).__name__}: {e}")
        return _empty_outputs(f"❌ Analysis failed / تعذّر التحليل: {type(e).__name__}")
    report.contract_type = _contract_type_label(contract_type, other_type, report.language)
    report.scope_notes = _scope_notes(contract_type, selected_laws or [], report.language)

    if report.total_clauses == 0:
        return _empty_outputs(
            "❌ لم يُتعرّف على أي بند في النص. تأكد/ي أن البنود مفصولة بأسطر وتبدأ بعناوين مثل \"البند الأول:\" أو \"المادة 1\".\n"
            "❌ No clauses were recognized. Make sure clauses are on separate lines and start with headings such as \"Article 1:\" or \"1.\"."
        )

    result = _summary_html(report, source_note, ml_pred)
    is_ar = (report.language == 'ar')

    fig = None
    charted = [lid for lid, s in report.law_scores.items() if s.compliance_pct is not None]
    if charted:
        laws_en = [LAWS[lid]['en'] for lid in charted]
        fig = go.Figure()
        fig.add_trace(go.Bar(name='Compliance %', x=laws_en,
                             y=[report.law_scores[lid].compliance_pct for lid in charted],
                             marker_color=STATUS_COLORS['compliance']))
        fig.add_trace(go.Bar(name='Risk % (severity)', x=laws_en,
                             y=[report.law_scores[lid].risk_score for lid in charted],
                             marker_color=STATUS_COLORS['risk']))
        fig.update_layout(barmode='group', height=360, template='plotly_white',
                          font=dict(family="IBM Plex Sans Arabic, Tahoma, sans-serif", color=TEXT),
                          legend=dict(orientation='h', y=-0.18, x=0.5, xanchor='center'),
                          yaxis=dict(range=[0, 100], gridcolor='#E6ECE9'), margin=dict(l=10, r=10, t=20, b=60))
    else:
        # No selected law has an applicable clause, so there is nothing to
        # plot; show the reason instead of leaving the chart area blank.
        fig = go.Figure()
        fig.add_annotation(
            text=("لا توجد بنود تتعلق بالأنظمة المختارة، لذلك لا يوجد ما يُعرض في الرسم"
                  if is_ar else
                  "No clause relates to the selected laws, so there is nothing to chart"),
            x=0.5, y=0.5, xref='paper', yref='paper', showarrow=False, font=dict(size=14))
        fig.update_layout(height=200, template='plotly_white',
                          xaxis=dict(visible=False), yaxis=dict(visible=False),
                          margin=dict(l=10, r=10, t=10, b=10))

    _cleanup_old_reports()
    pdf_path = save_pdf_report(report, str(_REPORT_DIR / f"report_{report.contract_id}.pdf"))

    return result, _trace_dataframe(report), fig, pdf_path


def _load_example(key):
    ex = EXAMPLE_CONTRACTS[key]
    # Also clear any uploaded file: a file takes precedence over the text box,
    # so leaving one attached would analyze the file instead of the example.
    return ex["name"], ex["text"], None, "services", ""


# ------------------------------------------------------------------
# Visual design: a sober ink-green palette suited to a legal/regulatory
# tool, one bilingual typeface (IBM Plex Sans Arabic, which covers Arabic
# and Latin), and status colours shared by the chart and the PDF report.
# ------------------------------------------------------------------
INK = "#0E4A3F"
INK_DARK = "#0A372F"
PAGE = "#F4F6F5"
TEXT = "#1C2622"
MUTED = "#5C6B65"
LINE = "#D6DEDA"
BRASS = "#9A7B3C"
STATUS_COLORS = {"compliance": "#2E7D5B", "risk": "#B3412F"}

_ink_hue = gr.themes.Color(
    c50="#EEF5F3", c100="#D5E7E2", c200="#ABCFC5", c300="#7DB3A5", c400="#4F9483",
    c500="#2A7666", c600="#155E51", c700="#0E4A3F", c800="#0A372F", c900="#072720",
    c950="#041813", name="ink")

_theme = gr.themes.Base(
    primary_hue=_ink_hue,
    neutral_hue=gr.themes.colors.gray,
    radius_size=gr.themes.sizes.radius_sm,
    font=[gr.themes.GoogleFont("IBM Plex Sans Arabic"), "Tahoma", "sans-serif"],
).set(
    body_background_fill=PAGE,
    body_text_color=TEXT,
    block_background_fill="#FFFFFF",
    block_border_color=LINE,
    block_border_width="1px",
    block_shadow="none",
    block_label_text_color=MUTED,
    block_title_text_color=TEXT,
    block_title_text_weight="600",
    input_background_fill="#FFFFFF",
    input_border_color=LINE,
    button_primary_background_fill=INK,
    button_primary_background_fill_hover=INK_DARK,
    button_primary_text_color="#FFFFFF",
    button_secondary_background_fill="#FFFFFF",
    button_secondary_border_color=LINE,
    button_secondary_text_color=INK,
    checkbox_background_color_selected=INK,
    checkbox_border_color_selected=INK,
    color_accent=INK,
    color_accent_soft="#EEF5F3",
)

_CSS = f"""
.gradio-container {{ max-width: 1180px !important; margin: 0 auto !important; }}
#app-header {{ background: {INK}; color: #fff; border-radius: 6px; padding: 26px 30px 22px; margin-bottom: 6px; }}
#app-header .row {{ display: flex; justify-content: space-between; align-items: flex-end; gap: 24px; flex-wrap: wrap; }}
#app-header h1 {{ font-size: 30px; font-weight: 600; margin: 0; color: #fff; line-height: 1.25; }}
#app-header .en {{ font-size: 17px; font-weight: 400; color: #CFE3DD; margin: 4px 0 0; }}
#app-header .meta {{ font-size: 13.5px; color: #CFE3DD; margin: 0; line-height: 1.7; max-width: 46ch; }}
#app-header .rule {{ height: 2px; width: 56px; background: {BRASS}; margin: 14px 0 0; }}
.step-title {{ font-weight: 600; font-size: 15px; color: {INK}; margin: 2px 0 -4px; }}
.step-title span {{ color: {MUTED}; font-weight: 400; }}
#analyze-btn {{ font-size: 16px; font-weight: 600; padding: 12px; }}
#results-box textarea {{ font-size: 14.5px; line-height: 1.75; }}
#trace-table table {{ font-size: 13.5px; }}
#feedback {{ border: 1px solid {LINE}; border-inline-start: 4px solid {BRASS}; background: #fff; border-radius: 6px; padding: 18px 22px; margin-top: 8px; }}
#feedback h3 {{ margin: 0 0 6px; color: {TEXT}; font-size: 17px; }}
#feedback p {{ margin: 4px 0; color: {TEXT}; line-height: 1.7; }}
#feedback a {{ color: {INK}; font-weight: 600; }}
footer {{ display: none !important; }}
#trace-table, #trace-table * {{ font-family: 'IBM Plex Sans Arabic', Tahoma, sans-serif !important; }}
#results-box {{ background: #fff; border: 1px solid {LINE}; border-radius: 6px; padding: 18px 20px; min-height: 120px; }}
.sum-empty {{ color: {MUTED}; text-align: center; padding: 40px 0; line-height: 2; }}
.sum-msg p {{ color: #B3412F; margin: 4px 0; font-weight: 500; }}
.sum-meta {{ color: {MUTED}; font-size: 13px; margin: 0 0 12px; }}
.sum-head {{ border-inline-start: 4px solid {LINE}; padding: 4px 14px; margin-bottom: 14px; }}
.sum-status {{ font-size: 21px; font-weight: 600; }}
.sum-figs {{ font-size: 15px; color: {TEXT}; margin-top: 4px; }}
.sum-basis {{ font-size: 13px; color: {MUTED}; margin-top: 2px; }}
.sum-table {{ width: 100%; border-collapse: collapse !important; font-size: 14px; border: none !important; }}
.sum-table th {{ text-align: start; color: {MUTED}; font-weight: 500; border: none !important; border-bottom: 1px solid {LINE} !important; padding: 6px 6px; background: transparent !important; }}
.sum-table td {{ border: none !important; border-bottom: 1px solid #EDF1EF !important; padding: 8px 6px; vertical-align: top; }}
.sum-table td.cnt {{ color: {MUTED}; font-size: 12.5px; }}
.sum-table .dot {{ display: inline-block; width: 9px; height: 9px; border-radius: 50%; margin-inline-end: 8px; }}
.sum-note {{ color: {MUTED}; font-size: 13px; margin: 12px 0 0; }}
.sum-warn {{ background: #FFF7E6; border: 1px solid #F0D9A8; color: #6B4A00; border-radius: 8px;
            padding: 8px 12px; font-size: 13.5px; margin: 10px 0 0; }}
#results-box table, #results-box thead, #results-box tbody, #results-box tr {{ border: none !important; box-shadow: none !important; outline: none !important; }}
#results-box tr > :first-child, #results-box tr > :last-child {{ border-left: none !important; border-right: none !important; }}
"""

_HEADER_HTML = """
<div id="app-header">
  <div class="row">
    <div dir="rtl" style="text-align:right">
      <h1>محلل امتثال العقود للأنظمة السعودية</h1>
      <p class="en" dir="ltr" style="text-align:right">Contract Compliance Analyzer for Saudi Regulations</p>
    </div>
    <div>
      <p class="meta" dir="rtl" style="text-align:right">نموذج أولي بحثي (المرحلة الأولى) لرسالة ماجستير في نظم المعلومات، جامعة الملك عبدالعزيز. يفحص بنود العقد مقابل مواد مدققة من ستة أنظمة سعودية، ويبيّن أساس كل حكم.</p>
      <div class="rule" style="margin-inline-start:auto"></div>
    </div>
  </div>
</div>
"""

_FEEDBACK_HTML = f"""
<div id="feedback">
  <div dir="rtl" style="text-align:right">
    <h3>تقييمك هو الأهم</h3>
    <p>بعد مراجعة النتائج وجدول تتبع البنود، يُرجى تعبئة <a href="{FEEDBACK_FORM_URL}" target="_blank" rel="noopener">نموذج التقييم</a>. تُستخدم الإجابات في فصل التقييم بالرسالة (الصحة، والوضوح، وقابلية التتبع، وسهولة الاستخدام).</p>
  </div>
  <div dir="ltr" style="text-align:left; margin-top:10px">
    <p>After reviewing the results and the clause traceability table, please complete the <a href="{FEEDBACK_FORM_URL}" target="_blank" rel="noopener">evaluation form</a>. Your answers feed the thesis evaluation chapter (correctness, clarity, traceability, usability).</p>
  </div>
</div>
"""


def _step(n: int, ar: str, en: str) -> str:
    return f'<p class="step-title" dir="rtl" style="text-align:right">{n}. {ar} <span>/ {en}</span></p>'


# Gradio deletes uploaded files and served copies of the PDF from its cache:
# checked every 15 minutes, removing files older than one hour.
_blocks_kwargs = dict(title="محلل امتثال العقود | Contract Compliance Analyzer",
                      theme=_theme, css=_CSS)
if "delete_cache" in inspect.signature(gr.Blocks.__init__).parameters:
    _blocks_kwargs["delete_cache"] = (900, 3600)

with gr.Blocks(**_blocks_kwargs) as demo:
    gr.HTML(_HEADER_HTML)
    with gr.Accordion("طريقة الاستخدام والقيود المعروفة / How to use and known limitations", open=False):
        with gr.Tabs():
            with gr.Tab("العربية"):
                gr.Markdown(NOTE_MD_AR, rtl=True)
            with gr.Tab("English"):
                gr.Markdown(NOTE_MD_EN)

    with gr.Row(equal_height=False):
        with gr.Column(scale=5):
            gr.HTML(_step(1, "العقد", "Contract"))
            with gr.Row():
                load_ar_btn = gr.Button("تحميل مثال عربي / Arabic example", size="sm")
                load_en_btn = gr.Button("تحميل مثال إنجليزي / English example", size="sm")
            contract_name = gr.Textbox(label="اسم العقد / Contract name", placeholder="مثال: عقد تقديم خدمات تقنية")
            contract_file = gr.File(label="رفع ملف PDF أو DOCX أو TXT (يُقدَّم على النص) / Upload PDF, DOCX or TXT (used instead of the text box)",
                                     file_types=[".pdf", ".docx", ".txt"], height=120)
            contract_text = gr.Textbox(label="أو ألصق/ي نص العقد / Or paste the contract text", lines=8)

            gr.HTML(_step(2, "نوع العقد والأنظمة", "Contract type and laws"))
            contract_type = gr.Dropdown(
                choices=[(f"{ar} / {en}", k) for k, ar, en in CONTRACT_TYPES],
                value=None,
                label="نوع العقد / Contract type",
            )
            other_type = gr.Textbox(label="حدّد/حدّدي نوع العقد / Specify the contract type",
                                    visible=False)
            selected_laws = gr.CheckboxGroup(
                choices=[(f"{LAWS[k]['ar']} / {LAWS[k]['en']}", k) for k in LAWS],
                value=[],
                label="الأنظمة التي تحكم عقدك / Laws that govern your contract"
            )

            gr.HTML(_step(3, "التحليل", "Analysis"))
            analyze_btn = gr.Button("تحليل العقد / Analyze contract", variant="primary", elem_id="analyze-btn")

        with gr.Column(scale=6):
            gr.HTML('<p class="step-title" dir="rtl" style="text-align:right">النتائج <span>/ Results</span></p>')
            result_text = gr.HTML(
                '<div class="sum-empty" dir="rtl">ستظهر النتائج هنا بعد التحليل.<br>'
                '<span dir="ltr">Results will appear here after the analysis.</span></div>',
                elem_id="results-box")
            result_chart = gr.Plot(label="التوافق والمخاطر لكل نظام / Compliance and risk by law")
            pdf_file = gr.File(label="تقرير PDF / PDF report")

    trace_table = gr.Dataframe(
        label="تتبع البنود: المادة المطابقة وأساس الحكم / Clause traceability: matched article and basis of each verdict",
        interactive=False, wrap=True, elem_id="trace-table",
    )

    gr.HTML(_FEEDBACK_HTML)

    load_ar_btn.click(lambda: _load_example("ar_tech_services"), None,
                      [contract_name, contract_text, contract_file, contract_type, other_type])
    load_en_btn.click(lambda: _load_example("en_service_agreement"), None,
                      [contract_name, contract_text, contract_file, contract_type, other_type])
    contract_type.change(lambda k: gr.update(visible=(k == "other")), contract_type, other_type)

    analyze_btn.click(
        analyze_ui,
        [contract_text, contract_name, contract_type, other_type, selected_laws, contract_file],
        [result_text, trace_table, result_chart, pdf_file]
    )

if __name__ == "__main__":
    demo.queue().launch()
